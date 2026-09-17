#!/usr/bin/env python
# -*- coding: utf-8 -*-
from __future__ import annotations

import json
import os
import queue
import re
import shutil
import subprocess
import threading
import time
import traceback
import urllib.request
import urllib.error
from urllib.parse import urlparse
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from pathlib import Path
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from course_state import (CourseStateStore, DEFAULT_TARGETS, completion_snapshot,
                          timetable_slots, layout_timetable_slots, teaching_class_core)

# ---- 高分屏 DPI 适配：必须在创建任何窗口前调用 ----
import ctypes


def _enable_dpi_awareness() -> float:
    """启用 Windows DPI 感知，返回系统缩放系数（1.0 = 100%）。"""
    try:
        ctypes.windll.shcore.SetProcessDpiAwareness(1)  # system DPI aware
    except Exception:
        try:
            ctypes.windll.user32.SetProcessDPIAware()
        except Exception:
            pass
    try:
        dpi = ctypes.windll.user32.GetDpiForSystem()
        return dpi / 96.0 if dpi > 0 else 1.0
    except Exception:
        return 1.0


DPI_SCALE = _enable_dpi_awareness()

BASE_URL_DEFAULT = "https://jwxt.zjgsu.edu.cn/jwglxt/"
COOKIE_DOMAIN = "jwxt.zjgsu.edu.cn"
PROJECT_DIR = Path(__file__).resolve().parent
# 资源与用户数据分离：正式版可整体替换升级，配置/状态/日志不丢失。
APP_NAME = "工商大学选课助手"
APP_VERSION = "v0.1.9"
DATA_DIR = Path(os.environ.get("LOCALAPPDATA", Path.home() / "AppData" / "Local")) / APP_NAME
DATA_LOG_DIR = DATA_DIR / "logs"
PRESET_FILE = DATA_DIR / "course_presets.txt"
PRESET_DETAILS_FILE = PROJECT_DIR / "course_preset_details.json"
SELECTED_STATE_FILE = DATA_DIR / "selected_courses.json"
SYSTEM_SELECTED_FILE = DATA_DIR / "system_selected_courses.json"
COURSE_CATALOG_FILE = DATA_DIR / "course_catalog.json"
SETTINGS_FILE = DATA_DIR / "settings.json"
def find_edge_exe() -> Path:
    """Find Microsoft Edge without depending on this app's absolute path."""
    candidates = []
    for env_name in ("PROGRAMFILES(X86)", "PROGRAMFILES", "LOCALAPPDATA"):
        root = os.environ.get(env_name)
        if root:
            candidates.append(Path(root) / "Microsoft" / "Edge" / "Application" / "msedge.exe")
    found = shutil.which("msedge.exe") or shutil.which("msedge")
    if found:
        candidates.insert(0, Path(found))
    return next((path for path in candidates if path.exists()), candidates[0] if candidates else Path("msedge.exe"))


EDGE_EXE = find_edge_exe()
EDGE_PROFILE = DATA_DIR / "edge_profile_zjgsu"
DEBUG_PORT = 9227
LOG_DIR = DATA_LOG_DIR
LOG_BUDGET = 1024 * 1024
LOG_ROTATE_SIZE = 256 * 1024

BG = "#f7f7f8"
PANEL = "#ffffff"
TEXT = "#202123"
MUTED = "#6b7280"
BORDER = "#d9d9e3"
SOFT = "#ececf1"
ACCENT = "#10a37f"
DANGER = "#dc2626"
WARN = "#b45309"


@dataclass
class QueryResult:
    client: object
    course: object
    jxbs: list
    keyword: str


@dataclass
class LocalSearchResult:
    rows: list
    keyword: str
    source: str = "本地缓存"
    hidden: int = 0      # 被校区筛选隐藏掉的个数（教工路校区的课）
    campus: str = ""     # 本次筛选用的校区
    weeks_hidden: int = 0    # 被周数筛选隐藏掉的个数（含时间/周次识别不出来的）
    weeks_cond: str = ""     # 本次周数筛选的输入（如 "8-13"）
    weeks_mode: str = ""     # 本次周数筛选的模式（"含于" / "剔除"）
    credits_hidden: int = 0  # 被学分筛选隐藏掉的个数（含学分未知的）
    credits_cond: str = ""   # 本次学分筛选选中的值（如 "1/2"）


def normalize_base(base: str) -> str:
    base = (base or BASE_URL_DEFAULT).strip()
    if not base.endswith("/"):
        base += "/"
    parsed = urlparse(base)
    if parsed.scheme != "https" or parsed.hostname != COOKIE_DOMAIN:
        raise ValueError("服务器地址必须是 https://jwxt.zjgsu.edu.cn/ 下的地址")
    if not parsed.path.rstrip("/").endswith("/jwglxt"):
        raise ValueError("服务器地址必须指向 /jwglxt/")
    return base


def mask_cookie(cookie: str) -> str:
    out = []
    for item in cookie.split(";"):
        item = item.strip()
        if not item:
            continue
        if "=" not in item:
            out.append(item[:3] + "***")
            continue
        k, v = item.split("=", 1)
        out.append(f"{k}={v[:4]}***{v[-3:] if len(v) > 7 else ''}")
    return "; ".join(out)


def build_edge_args(base_url: str, port: int = DEBUG_PORT):
    return [
        str(EDGE_EXE),
        f"--remote-debugging-port={port}",
        "--remote-debugging-address=127.0.0.1",
        f"--remote-allow-origins=http://127.0.0.1:{port}",
        f"--user-data-dir={EDGE_PROFILE}",
        "--new-window",
        normalize_base(base_url),
    ]


def launch_login_browser(base_url: str, log):
    if not EDGE_EXE.exists():
        raise RuntimeError("没有找到 Microsoft Edge。")
    EDGE_PROFILE.mkdir(parents=True, exist_ok=True)
    subprocess.Popen(build_edge_args(base_url), stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    log("已打开独立 Edge 登录窗口。登录成功后回到这里点“自动抓取 Cookie”。")


def close_login_browser(log):
    import websocket
    tabs = _json_get(f"http://127.0.0.1:{DEBUG_PORT}/json")
    target = next((tab for tab in tabs if tab.get("webSocketDebuggerUrl")), None)
    if not target:
        raise RuntimeError("没有找到独立 Edge 调试窗口。")
    ws = websocket.create_connection(target["webSocketDebuggerUrl"], timeout=5)
    try:
        ws.send(json.dumps({"id": 99, "method": "Browser.close"}))
    finally:
        ws.close()
    log("已关闭独立 Edge，调试端口已释放。")


def _json_get(url: str, timeout=3):
    with urllib.request.urlopen(url, timeout=timeout) as r:
        return json.loads(r.read().decode("utf-8", "replace"))


def capture_cookie_from_edge(log) -> str:
    try:
        tabs = _json_get(f"http://127.0.0.1:{DEBUG_PORT}/json")
    except Exception as e:
        raise RuntimeError("没有连上登录浏览器。请先点“打开登录窗口”。") from e
    candidates = [tab for tab in tabs if COOKIE_DOMAIN in tab.get("url", "")]
    # 同时打开门户、首页和自主选课页时，优先绑定当前带 CAS ticket 的选课页。
    # 直接取第一个标签页可能拿到未完成 SSO 的首页，导致 Cookie 看似完整但无法访问选课入口。
    target = max(candidates, key=lambda tab: (
        "/xsxk/" in tab.get("url", ""),
        "ticket=" in tab.get("url", ""),
        tab.get("type") == "page",
    ), default=None)

    if not target and tabs:
        target = tabs[0]
    if not target or not target.get("webSocketDebuggerUrl"):
        raise RuntimeError("没找到可读取的浏览器标签页。")

    import websocket
    try:
        ws = websocket.create_connection(target["webSocketDebuggerUrl"], timeout=5)
    except Exception as e:
        raise RuntimeError("浏览器拒绝 WebSocket 调试连接。已修复启动参数，请关闭旧的独立 Edge 登录窗口后，重新点“打开登录窗口”。") from e
    try:
        ws.send(json.dumps({"id": 1, "method": "Network.enable"}))
        ws.recv()
        ws.send(json.dumps({"id": 2, "method": "Network.getCookies", "params": {
            "urls": ["https://jwxt.zjgsu.edu.cn/jwglxt/xsxk/zzxkyzb_cxZzxkYzbIndex.html"]
        }}))
        cookies = []
        deadline = time.time() + 5
        while time.time() < deadline:
            msg = json.loads(ws.recv())
            if msg.get("id") == 2:
                cookies = msg.get("result", {}).get("cookies", [])
                break
    finally:
        ws.close()

    pairs = []
    for c in cookies:
        d, n, v = c.get("domain", ""), c.get("name", ""), c.get("value", "")
        if n and ("zjgsu.edu.cn" in d or d.endswith(".zjgsu.edu.cn")):
            pairs.append(f"{n}={v}")
    if not pairs:
        raise RuntimeError("没有抓到浙江工商教务 Cookie，请确认独立 Edge 已登录成功。")
    cookie = "; ".join(pairs)
    log(f"已抓取 {len(pairs)} 条 Cookie：{mask_cookie(cookie)}")
    return cookie


def parse_cookie_pairs(cookie: str) -> dict:
    pairs = {}
    for item in cookie.split(";"):
        item = item.strip()
        if not item or "=" not in item:
            continue
        k, v = item.split("=", 1)
        if k.strip():
            pairs[k.strip()] = v.strip()
    return pairs


def validate_cookie_format(cookie: str) -> tuple[bool, str]:
    pairs = parse_cookie_pairs(cookie)
    if not pairs:
        return False, "Cookie 为空或格式不对，应该类似 JSESSIONID=xxx; route=xxx"
    has_session = any(k.upper() in {"JSESSIONID", "SESSION", "JSESSIONIDSSO"} or "SESSION" in k.upper() for k in pairs)
    if not has_session:
        return False, "Cookie 里没有明显的 Session 字段，可能复制错页面了。"
    return True, f"Cookie 格式正常，共 {len(pairs)} 项：" + ", ".join(list(pairs.keys())[:8])


def http_verify_cookie(base_url: str, cookie: str) -> tuple[str, str]:
    url = normalize_base(base_url) + "xtgl/index_initMenu.html"
    req = urllib.request.Request(url, headers={
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/120 Safari/537.36 Edg/120",
        "Cookie": cookie,
        "Referer": normalize_base(base_url),
    })
    opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor())
    try:
        resp = opener.open(req, timeout=10)
        final_url = resp.geturl()
        body = resp.read(5000).decode("utf-8", "replace")
    except urllib.error.HTTPError as e:
        final_url = e.geturl()
        body = e.read(3000).decode("utf-8", "replace")
    except Exception as e:
        return "unknown", f"HTTP 验证失败：{type(e).__name__}: {e}"

    low = (final_url + "\n" + body).lower()
    if "cas/login" in low or "统一身份认证" in body or "login" in final_url.lower():
        return "invalid", "HTTP 验证：被跳回统一认证/登录页，Cookie 大概率无效或不完整。"
    if "sessionuserkey" in low or "index_initmenu" in low or "教学管理" in body or "jwglxt" in final_url.lower():
        return "valid", "HTTP 验证：能访问教务系统页面，Cookie 基本有效。"
    return "unknown", f"HTTP 验证：未能明确判断。最终地址：{final_url}"


def verify_cookie_full(base_url: str, cookie: str, log) -> str:
    ok, msg = validate_cookie_format(cookie)
    log("[格式检查] " + msg)
    if not ok:
        return "Cookie格式错误"
    http_status, http_msg = http_verify_cookie(base_url, cookie)
    log("[HTTP检查] " + http_msg)
    client = make_client(base_url, cookie, log)
    try:
        key = client.check_login()
        if key:
            log("[登录检查] check_login 成功：已登录")
            return "Cookie有效，且已登录"
        log("[登录检查] check_login 未通过（可能已过期）")
    except Exception as e:
        log("[登录检查] check_login 失败：" + type(e).__name__ + ": " + str(e))
    try:
        ok_init = client.init()
        if ok_init:
            log("[选课检查] init 成功：选课已开放/参数可用")
            return "Cookie有效，选课已开放"
        log("[选课检查] init 未成功：选课入口拿不到参数")
    except Exception as e:
        log("[选课检查] init 未成功：" + type(e).__name__ + ": " + str(e))
    if http_status == "valid":
        return "Cookie基本有效，等待选课开放"
    if http_status == "invalid":
        return "Cookie无效或已过期"
    return "无法确认Cookie，请重新登录后抓取"


def make_client(base_url: str, cookie: str, log):
    from zjgsu_api import ZjgsuClient
    client = ZjgsuClient(cookie, base_url=base_url, log=log)
    log("已使用浙工商直连模块（requests 复刻，绕过 lnuElytra 编译绑定）。")
    return client


def init_until_open(client, log, once=False) -> bool:
    try:
        ok = bool(client.init())
    except Exception as e:
        log(f"init 未成功：{type(e).__name__}: {e}")
        ok = False
    if ok:
        log("init 成功：选课入口/参数已可用。")
        return True
    log("init 未成功：选课入口拿不到参数（可能未开放或 Cookie 失效）。")
    if once:
        raise RuntimeError("选课尚未开放：init 未成功")
    return False


def normalize_mode(mode: str) -> str:
    return {
        "自动识别": "auto",
        "课程代码": "course_code",
        "教学班": "teaching_class",
        "课程名": "course_name",
        "auto": "auto",
        "course_code": "course_code",
        "teaching_class": "teaching_class",
        "course_name": "course_name",
    }.get(mode or "自动识别", "auto")


def normalize_keyword(keyword: str, mode: str = "自动识别") -> str:
    kw = (keyword or "").strip()
    mode_key = normalize_mode(mode)
    # lnuElytra/正方接口本质上把这里作为 filter_list[0] 发给服务器。
    # 课程代码、教学班号、课程名都可以作为同一个 q 使用；这里保留模式是为了 UI 表达更清楚。
    return kw


def fetch_course(client, keyword: str, mode: str = "auto") -> QueryResult:
    q = normalize_keyword(keyword, mode)
    course = client.fetch_course(q) if hasattr(client, "fetch_course") else client.fetch_courses(q)
    return QueryResult(client=client, course=course, jxbs=list(getattr(course, "jxb", []) or []), keyword=q)


def choose_jxb(jxbs, selector: str):
    selector = (selector or "").strip()
    if not jxbs:
        return None
    if not selector:
        return jxbs[0]
    for j in jxbs:
        hay = " | ".join(str(getattr(j, name, "") or "") for name in ("sksj", "jsxx", "jxb_id", "do_id"))
        if selector in hay:
            return j
    return None


def select_jxb(client, course, jxb, log):
    kch_id = getattr(course, "kch_id", None)
    do_id = getattr(jxb, "do_id", None)
    xkkz_id = getattr(course, "xkkz_id", None) or None
    log(f"提交：kch_id={kch_id}, do_id={do_id[:16] if do_id else None}..., xkkz_id={str(xkkz_id)[:8] if xkkz_id else None}")
    try:
        res = client.select_course(kch_id, do_id, xkkz_id=xkkz_id)
    except TypeError:
        res = client.select_course(kch_id, do_id)
    # 直连模块返回 (flag, msg) 元组；lnuElytra pyd 返回对象。
    if isinstance(res, tuple):
        flag, msg = res
    else:
        flag = getattr(res, "flag", None)
        msg = str(getattr(res, "msg", None) or "")
    return flag, msg


QUOTA_MARKER_RE = re.compile(r"\b-?\d+,[0-9A-Za-z]{16,},\d+\b")


def is_quota_marker(msg) -> bool:
    """识别教务“没有中文提示”的名额返回，例如 "0,520F45EC...,33"。

    保底志愿轮次对满课教学班就返回这种串。旧代码把它当普通失败原样打印，
    看起来像“指令没发出去”，实际上提交已经到达服务器。识别出来后按满课处理：
    保留 do_id 缓存、节流盲提交蹲退课窗口。
    """
    return bool(QUOTA_MARKER_RE.search(str(msg or "")))


def annotate_select_msg(msg: str) -> str:
    """给难以解读的教务返回加上人话前缀，日志里能一眼看懂。"""
    text = str(msg or "").strip()
    if not text:
        return text
    if is_quota_marker(text) and "名额标记" not in text:
        return f"教务名额标记 {text}（疑似已满/保底志愿轮次，按未成功继续蹲守）"
    return text


def describe_select_result(flag, msg) -> str:
    """把一次提交的 flag/msg 翻译成一句可直接读的结论。"""
    if str(flag) == "1":
        return f"提交成功（flag=1）：{msg}" if msg else "提交成功（flag=1）"
    text = str(msg or "").strip()
    if not text:
        return f"未成功：教务未返回提示（flag={flag}）"
    if is_quota_marker(text):
        return f"未成功：教务名额标记 {text}（疑似已满/保底志愿，程序按未成功继续尝试）"
    return f"未成功：{text}"


def should_continue_for_msg(msg: str, retry_rate: bool, watch_full: bool) -> tuple[bool, float, str]:
    if "频率过高" in msg and retry_rate:
        return True, 1.0, "频率过高，1 秒后重试"
    if watch_full and (any(x in msg for x in ["容量已满", "已满", "满员"]) or is_quota_marker(msg)):
        return True, 5.0, "容量已满，蹲守退课"
    if "未开放" in msg:
        return True, 0.0, "未开放，继续监控"
    return False, 0.0, ""


def pull_all_courses(base_url: str, cookie: str, log):
    """用直连模块分页拉取全部教学班，并并行补上时间/教室/教师/人数等详情。"""
    from zjgsu_api import ZjgsuClient
    client = ZjgsuClient(cookie, base_url=base_url, log=log)
    if not client.init():
        raise RuntimeError("选课尚未开放：选课入口拿不到参数，请等开放后再试。")
    rows = client.fetch_all_courses_detailed(page_size=1000, max_pages=200, workers=8, log=log)
    if not rows:
        raise RuntimeError("没有拉到任何课程，请检查 Cookie 是否过期。")
    out = []
    for r in rows:
        out.append({
            "keyword": r.get("kcmc", ""),          # 课程名（兼容旧字段）
            "kch_id": r.get("kch_id", ""),
            "jxb_count": 1,                         # 每行=一个教学班
            "course_name": r.get("kcmc", ""),
            "kch": r.get("kch", ""),                # 课程号
            "teaching_class": r.get("jxbmc", ""),   # 教学班名称
            "xf": r.get("xf", ""),                  # 学分
            "teacher": r.get("jsxx", ""),           # 教师
            "time": r.get("sksj", ""),              # 上课时间
            "room": r.get("jxdd", ""),              # 上课教室
            "yixuan": r.get("yxzrs", ""),           # 已选人数
            "rongliang": r.get("jxbrl", ""),        # 容量
            "xqumc": r.get("xqumc", ""),             # 上课校区（教务权威字段）
            "xueyuan": r.get("kkxymc", ""),         # 开课学院
            "kklxdm": (r.get("_course_tab") or {}).get("kklxdm", ""),
            "course_tab": (r.get("_course_tab") or {}).get("name", ""),
            "course_belong": r.get("course_belong", "") or r.get("kccat", ""),
            "course_group": r.get("course_group", "") or r.get("kclbmc", ""),
            "kclbmc": r.get("kclbmc", ""),
            "kcxzmc": r.get("kcxzmc", ""),
            "do_id": r.get("do_id", ""),
            "jxb_id": r.get("jxb_id", ""),
        })
    log(f"[全部课程] 直连拉取完成：{len(out)} 个教学班（含时间/教室详情）")
    return out


def search_cached_courses(rows, keyword, mode="自动识别", filters=None):
    """在已拉取的全部课程缓存中本地搜索，不依赖选课窗口是否开放。"""
    needle = (keyword or "").strip().casefold()
    if not needle and not filters:
        return []
    fields = {
        "课程代码": ("kch", "course_code"),
        "教学班": ("teaching_class", "jxb_id"),
        "课程名": ("course_name", "keyword"),
    }.get(mode, ("course_name", "course_code", "teaching_class", "kch", "jxb_id", "teacher"))
    filters = filters or {}
    selected_categories = set(filters.get("course_categories", filters.get("course_tabs", [])))
    selected_belongs = set(filters.get("course_belongs", []))
    selected_groups = set(filters.get("course_groups", []))
    selected_days = set(filters.get("days", []))
    selected_periods = set(str(x) for x in filters.get("periods", []))
    availability = filters.get("availability", "全部")
    selected_campus = filters.get("campus", CAMPUS_ALL)
    weeks_text = str(filters.get("weeks") or "").strip()
    wanted_weeks = week_set(weeks_text) if weeks_text else None
    weeks_mode = filters.get("weeks_mode") or WEEK_MODE_INCLUDE
    selected_credits = set(str(x) for x in (filters.get("credits") or []))
    result = []
    for row in rows or []:
        haystack = " ".join(str(row.get(field, "") or "") for field in fields).casefold()
        if needle not in haystack:
            continue
        category = course_category(row)
        belong = str(row.get("course_belong") or "").strip()
        group = str(row.get("course_group") or row.get("kclbmc") or "未分组")
        group_context = {group, str(row.get("course_tab") or "")}
        time_text = str(row.get("time") or row.get("sksj") or "")
        if selected_categories and category not in selected_categories:
            continue
        if selected_belongs:
            # 课程归属是"通识"类专属细分维度：只在课程属于通识时按归属过滤，
            # 非通识课程（主修/体育/英语）不受归属勾选影响，避免组合筛选被误杀。
            if category == "通识" and belong not in selected_belongs:
                continue
        if selected_groups and not any(course_group_matches(value, selected_groups) for value in group_context):
            continue
        if selected_days and not any(f"星期{day}" in time_text for day in selected_days):
            continue
        if selected_periods:
            occupied = occupied_periods(time_text)
            if not occupied.intersection(selected_periods):
                continue
        if not campus_filter_match(row, selected_campus):
            continue
        # 周数筛选：教务的课常只上其中几周（{8-13周}），同一时间段在不同周次里不真的撞车，
        # 所以单独按「含于 / 剔除」收窄（week_filter_match 的语义见其 docstring）。
        if wanted_weeks and not week_filter_match(time_text, wanted_weeks, weeks_mode):
            continue
        # 学分筛选：挑通识时最常用（通识选修课只有 1 分 / 2 分两档），
        # 学分取不到的课（缓存里没有 xf）在启用学分筛选时一并排除，并如实报数。
        if selected_credits and not credit_matches(row, selected_credits):
            continue
        try:
            remaining = int(float(row.get("rongliang", 0) or 0)) - int(float(row.get("yixuan", 0) or 0))
        except (TypeError, ValueError):
            remaining = None
        if availability == "有余量" and not (remaining is not None and remaining > 0):
            continue
        if availability == "无余量" and not (remaining is not None and remaining <= 0):
            continue
        item = dict(row)
        item["course_category"] = category
        item["course_belong"] = belong
        item["course_group"] = group
        item["campus"] = row_campus(row)
        result.append(item)
    return result


COURSE_GROUP_ALIASES = {
    "主修": {"主修", "主修课程", "专业课程", "专业必修课", "专业选修课"},
    "体育": {"体育", "体育课程", "体育分项", "体育课"},
    "英语": {"英语", "英语课程", "大学英语", "英语分项"},
    "思政/公共基础": {"思政", "思想政治理论课", "公共基础课", "公共基础", "军事理论", "形势与政策"},
    "通识/公共选修": {"通识", "通识选修课", "通识选修课课组", "通识教育", "公共艺术课", "五史课", "公共选修课"},
    "其他": {"其他", "其它", "未分组"},
}


COURSE_CATEGORY_LABELS = ("主修", "思想政治", "体育", "英语", "通识")


def course_category(row: dict) -> str:
    """Map the university course-tab fields to the five top-level categories."""
    code = str(row.get("kklxdm") or row.get("course_category_code") or "").strip()
    name = str(row.get("course_tab") or row.get("course_category") or "").strip()
    if code == "05" or "体育" in name:
        return "体育"
    if code == "07" or "英语" in name:
        return "英语"
    if code in {"18", "34"} or any(x in name for x in ("思政", "思想政治", "形势与政策")):
        return "思想政治"
    if code == "10" or "通识" in name:
        return "通识"
    if code == "01" or "主修" in name:
        return "主修"
    return "通识" if str(row.get("course_belong") or "").strip() else "主修"


def course_group_matches(actual: str, selected: set[str]) -> bool:
    actual = str(actual or "未分组").strip()
    for label in selected:
        if actual == label:
            return True
        if actual in COURSE_GROUP_ALIASES.get(label, set()):
            return True
    return False


def occupied_periods(time_text: str) -> set[str]:
    """Expand every lesson interval, so selecting one period matches overlaps."""
    occupied = set()
    for start, end in re.findall(r"第(\d+)(?:-(\d+))?节", str(time_text or "")):
        first = int(start)
        last = int(end or start)
        if last < first:
            first, last = last, first
        occupied.update(str(period) for period in range(first, last + 1))
    return occupied


def seconds_until_start(start_text: str) -> float:
    txt = (start_text or "").strip()
    if not txt:
        return 0.0
    from datetime import datetime, timedelta
    now = datetime.now()
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            t = datetime.strptime(txt, fmt).time()
            target = datetime.combine(now.date(), t)
            if target < now:
                target = target + timedelta(days=1)
            return max(0.0, (target - now).total_seconds())
        except ValueError:
            pass
    raise RuntimeError("定时开始格式错误，请填 HH:MM 或 HH:MM:SS，例如 09:59:58")


def timed_monitor_rush(start_text, *args, **kwargs):
    log = args[-1] if args else kwargs.get("log")
    wait = seconds_until_start(start_text)
    if wait > 0:
        if log:
            log(f"定时启动已设置：{start_text}，将在 {wait:.1f} 秒后开始监控。")
        stop_event = args[-2] if len(args) >= 2 else kwargs.get("stop_event")
        while wait > 0:
            if stop_event and stop_event.is_set():
                if log:
                    log("定时等待已取消。")
                return
            step = min(1.0, wait)
            time.sleep(step)
            wait -= step
    if log:
        log("进入监控抢课。")
    return monitor_rush(*args, **kwargs)


# 满课蹲守盲打节流：太快会被教务判"操作频率过高"，太慢又抢不到退课空位。
DEFAULT_FULL_WAIT = 2.5
FULL_WAIT_MIN, FULL_WAIT_MAX = 0.5, 10.0
# 「优先盯前 N 门」时，其余保底项每几轮试一次（1 轮 = 前 N 门各提交一次）
REST_EVERY_ROUNDS = 3

TEACHING_CLASS_RE = re.compile(r"^(?:[^\d(（】]{1,12})?\(\d{4}-\d{4}-\d\)-[A-Za-z0-9]+-[A-Za-z0-9-]+$")


def looks_like_teaching_class(value) -> bool:
    """判断一个预设关键字是教学班号（(2026-2027-1)-GENNET010-2）还是课程名。

    教务对非本校区的课会在教学班名前加校区前缀（教工路(2026-2027-1)-GENLHP076-03），
    这种同样算"锁定教学班"，否则会被误判成课程名、显示成"按课程名抢"。
    """
    return bool(TEACHING_CLASS_RE.match(str(value or "").strip()))


# ---------------------------------------------------------------------------
# 校区识别（下沙 = 我所在校区）
# ---------------------------------------------------------------------------
# 浙江工商大学：下沙校区（本部，我在这里上课）+ 教工路校区（杭州市区）。
# 教务不会给"是不是我的校区"这个字段，只能从三个地方推：
#   1) QueryDo 返回的 xqumc（教务权威字段，最准）；
#   2) 教学班名前缀：教务对非本校区课程加校区名（"教工路(2026-2027-1)-GENLHP076-03"）。
#      ⚠ 前缀不一定是校区！教务也给网络课加平台名（"智慧树网络课(...)"/"尔雅网络课(...)"），
#      所以只有「已知校区名」或以「校区」结尾的前缀才算校区，其余当前缀信息展示（网课）。
#   3) 教室名（不可靠，两个校区都有"5号教学楼"，只有明确写着校区名时才采信）。
CAMPUS_LOCAL = "下沙"
CAMPUS_ALL = "全部"
CAMPUS_NAMES = ("下沙", "教工路")
CAMPUS_ROOM_HINTS = {"教工路": ("教工路",)}
_TEACHING_CLASS_PREFIX_RE = re.compile(r"^\s*([^\d(（】]{1,12}?)\s*[(（]\s*\d{4}\s*[-—–－]\s*\d{4}")


def normalize_campus_name(value) -> str:
    """把 "教工路校区" / "教工路" 统一成 "教工路"；空值返回空串。"""
    text = str(value or "").strip()
    for suffix in ("校区", "校區", "区"):
        if text.endswith(suffix) and len(text) > len(suffix):
            text = text[: -len(suffix)].strip()
    return text


def is_campus_name(value) -> bool:
    """这个前缀/字段是不是校区名（而不是"智慧树网络课"这种授课平台）。"""
    text = str(value or "").strip()
    if not text:
        return False
    if normalize_campus_name(text) in CAMPUS_NAMES:
        return True
    return text.endswith(("校区", "校區"))


def teaching_class_prefix(teaching_class) -> str:
    """取教学班名/号里 "(学年-学期)" 之前的那段前缀（可能是校区，也可能是网课平台）。"""
    text = re.sub(r"^[^】]{0,40}】", "", str(teaching_class or "")).strip()
    match = _TEACHING_CLASS_PREFIX_RE.match(text)
    return match.group(1).strip() if match else ""


def teaching_class_campus(teaching_class) -> str:
    """教学班名前缀里的校区；不是校区（网课平台/课程别名）或没有前缀时返回空串。

    "教工路(2026-2027-1)-GENLHP076-03"        → "教工路"
    "教工路美术鉴赏(2026-2027-1)-GENARC033-01" → "教工路"（前缀里带课程别名）
    "(2026-2027-1)-GENNET045-2"               → ""
    "智慧树网络课(2026-2027-1)-X-1"            → ""（是网课平台，不是校区）
    "影视鉴赏(2026-2027-1)-GENARC012-03"       → ""（是课程别名，不是校区）
    """
    prefix = teaching_class_prefix(teaching_class)
    if not prefix:
        return ""
    if is_campus_name(prefix):
        return normalize_campus_name(prefix)
    # 前缀可能把校区名和课程别名连在一起（"教工路美术鉴赏"、"教工路创新研讨"）：
    # 只要以已知校区名开头就认这个校区，否则（"影视鉴赏"/"音乐鉴赏"…）不是校区。
    for name in CAMPUS_NAMES:
        if prefix.startswith(name):
            return name
    return ""


# 只有这些字样才说明是网络课（前缀也可能是课程别名，例如"影视鉴赏""美术鉴赏"，
# 那些既不是校区也不是网课，不能标成网课骗人）。
PLATFORM_PREFIX_RE = re.compile(r"网络课|网课|在线|慕课|MOOC|智慧树|尔雅|超星|学习通|学堂在线", re.I)


def teaching_class_platform(teaching_class) -> str:
    """非校区、非别名的教学班名前缀里，属于网课平台的那部分（否则返回空串）。"""
    prefix = teaching_class_prefix(teaching_class)
    if not prefix or teaching_class_campus(teaching_class):
        return ""
    return prefix if PLATFORM_PREFIX_RE.search(prefix) else ""


def detect_campus(xqumc="", teaching_class="", room="") -> str:
    """判定教学班所在校区（只用于标注/筛选，不参与提交）。

    两个来源只要有一个说"不是本校"，就按非本校处理（宁可多标一个 ⚠，
    也不要让下沙的学生抢到教工路的一门课）；但"网课平台"这种前缀不算校区。
    """
    candidates = (normalize_campus_name(xqumc), teaching_class_campus(teaching_class))
    for candidate in candidates:
        if candidate and candidate != CAMPUS_LOCAL:
            return candidate
    for candidate in candidates:
        if candidate:
            return candidate
    room_text = str(room or "")
    for name, hints in CAMPUS_ROOM_HINTS.items():
        if name in room_text or any(hint in room_text for hint in hints):
            return name
    return CAMPUS_LOCAL


def is_off_campus(campus) -> bool:
    """是否不在我所在校区（下沙）。空值按本校处理。"""
    name = normalize_campus_name(campus)
    return bool(name) and name != CAMPUS_LOCAL


def row_campus(row: dict) -> str:
    """从一行课程数据（缓存行/搜索结果/详情）里取校区。"""
    row = row or {}
    return detect_campus(row.get("xqumc") or row.get("campus") or "",
                         row.get("teaching_class") or row.get("jxbmc") or "",
                         row.get("room") or row.get("jxdd") or "")


def row_platform(row: dict) -> str:
    """从一行课程数据里取授课平台标记（网络课）；不是网课时返回空串。"""
    row = row or {}
    return row.get("platform") or teaching_class_platform(
        row.get("teaching_class") or row.get("jxbmc") or "")


def campus_label(campus) -> str:
    """界面上显示的校区标签。"""
    name = normalize_campus_name(campus) or CAMPUS_LOCAL
    return f"{name}（本校）" if name == CAMPUS_LOCAL else f"{name}（非本校）"


def campus_filter_match(row, selected) -> bool:
    """校区筛选：selected 为 "全部" 或空时全部通过，否则按校区名精确比对。

    selected 允许是列表（早期设置里误存成 ["下沙"] 的形状也要能用）。
    """
    if isinstance(selected, (list, tuple, set)):
        selected = next(iter(selected), CAMPUS_ALL)
    wanted = normalize_campus_name(selected)
    if not wanted or selected in (None, "", CAMPUS_ALL):
        return True
    return normalize_campus_name(row_campus(row)) == wanted


CATEGORY_PREFIX = {"GE": "通识", "EN": "英语", "PE": "体育"}


def week_numbers(weeks_text) -> set:
    """解析 "{1-16周}" / "{8-12周}" / "{1-5周,7周}" 里的周次集合。

    解析不出任何周次时返回 None，表示"按全周处理"（宁可当冲突，也不漏判）。
    """
    weeks = set()
    for part in re.split(r"[;,、，]", str(weeks_text or "")):
        match = re.search(r"(\d+)\s*(?:-\s*(\d+))?\s*周?", part)
        if not match:
            continue
        first = int(match.group(1))
        last = int(match.group(2) or first)
        weeks.update(range(min(first, last), max(first, last) + 1))
    return weeks or None


# 周数筛选的两种模式（搜索抢课 Tab「周数筛选」行）：
#   含于 —— 只保留上课周次**全部落在**输入区间内的课（输入 1-7：留 1-3、2-7，去 1-16、6-9）
#   剔除 —— 去掉任何一段与输入区间**有重叠周次**的课（输入 8-13：留 1-7、14-16）
WEEK_MODE_INCLUDE = "含于"
WEEK_MODE_EXCLUDE = "剔除"
WEEK_MODE_BUTTON_TEXT = {WEEK_MODE_INCLUDE: "只留含于", WEEK_MODE_EXCLUDE: "剔除重叠"}


def week_set(weeks_text) -> set:
    """解析 "{...}" 里的周次文本 → 周次集合（识别 "(单)/(双)" 修饰）。

    与 week_numbers 的分工：week_numbers 用于提交前的冲突判定，忽略单双（宁可当冲突）；
    本函数用于「周数筛选」，要如实 —— "1-16周(双)"只算 2/4/…/16 周。
    解析不出任何周次时返回 None（无法判定），由调用方决定怎么处理。
    """
    weeks = set()
    for part in re.split(r"[;,、，\s]+", str(weeks_text or "")):
        match = re.search(r"(\d+)\s*(?:-\s*(\d+))?\s*周?", part)
        if not match:
            continue
        first = int(match.group(1))
        last = int(match.group(2) or first)
        span = set(range(min(first, last), max(first, last) + 1))
        if "双" in part:
            span = {week for week in span if week % 2 == 0}
        elif "单" in part:
            span = {week for week in span if week % 2 == 1}
        weeks.update(span)
    return weeks or None


def course_week_segments(time_text) -> list:
    """把上课时间拆成每一段的周次集合；解析不出时段结构（如时间字段是 "--"）时返回 []。"""
    return [week_set(slot["weeks"]) for slot in timetable_slots([{"time": time_text}])]


def week_filter_match(time_text, wanted, mode: str = WEEK_MODE_INCLUDE) -> bool:
    """周数筛选：这门课的上课周次是否满足「含于 / 剔除」条件。

    时间或周次识别不出来的课（例如时间字段是 "--"、网课待定）两种模式下都不通过 ——
    筛完的列表只留"能确认满足条件"的行，被筛掉几条会在状态栏如实报出来。
    """
    wanted = set(wanted or ())
    if not wanted:
        return True
    segments = course_week_segments(time_text)
    if not segments:
        return False
    for weeks in segments:
        if weeks is None:
            return False
        if mode == WEEK_MODE_EXCLUDE:
            if weeks & wanted:
                return False
        elif not weeks <= wanted:
            return False
    return True


def credit_value(row) -> str:
    """取一门课的学分，规范化成显示字符串（"2.0" → "2"，"0.5" 保持）。

    学分只有一个来源：教务查询返回行里的 xf（缓存 course_catalog.json 的 xf 字段）。
    取不到时返回空串，界面上显示 "—"——不要假装成 0 分。
    """
    raw = row.get("xf") if isinstance(row, dict) else row
    if raw in (None, ""):
        return ""
    try:
        value = float(raw)
    except (TypeError, ValueError):
        return str(raw).strip()
    if value == int(value):
        return str(int(value))
    return "%g" % value


def credit_matches(row, selected) -> bool:
    """学分筛选：选中的学分集合是否包含这门课的学分（集合为空 = 不筛）。"""
    wanted = set(str(x) for x in (selected or ()))
    if not wanted:
        return True
    return credit_value(row) in wanted


def time_slots(time_text) -> list:
    """把上课时间文本拆成 [(星期, 起始节, 结束节, 周次集合)]，用于冲突判断。"""
    return [(slot["day"], slot["start"], slot["end"], week_numbers(slot["weeks"]))
            for slot in timetable_slots([{"time": time_text}])]


def time_conflicts(time_a, time_b, week_strict: bool = False) -> bool:
    """两段上课时间是否冲突（同一天 + 节次重叠）。

    week_strict=False（默认）：只要节次重叠就判冲突——教务的冲突提示多数按节次算，
    保守一点可以避免白提交。week_strict=True 时还要求周次有交集。
    """
    for day_a, start_a, end_a, weeks_a in time_slots(time_a):
        for day_b, start_b, end_b, weeks_b in time_slots(time_b):
            if day_a != day_b or start_a > end_b or start_b > end_a:
                continue
            if week_strict and weeks_a and weeks_b and not (weeks_a & weeks_b):
                continue
            return True
    return False


def catalog_matches(catalog: dict, keyword: str, category: str) -> list:
    """在本地「全部课程」缓存里找出某个关键字对应的教学班（可多个）。

    优先级：教学班号精确命中 → 课程号/课程名精确命中；同名多个时优先同类别 Tab
    （GE 优先通识选修课、EN 优先英语分项、PE 优先体育分项）。
    """
    catalog = catalog or {}
    hit = catalog.get(keyword)
    if hit:
        return [(keyword, hit)]
    wanted = str(keyword or "").strip().casefold()
    if not wanted:
        return []
    matches = [(teaching_class, row) for teaching_class, row in catalog.items()
               if wanted in {str(row.get("course_code") or "").strip().casefold(),
                             str(row.get("course_name") or "").strip().casefold()}]
    if not matches:
        return []
    preferred = CATEGORY_PREFIX.get(category)
    if preferred:
        same_tab = [(tc, row) for tc, row in matches
                    if preferred in str(row.get("course_tab") or "")
                    or preferred in str(row.get("kclbmc") or "")]
        if same_tab:
            matches = same_tab
    return matches


def capacity_room(row) -> tuple:
    """(是否还有余量, 剩余名额) —— 解析失败时按 0 处理。"""
    try:
        cap = int(float(row.get("rongliang") or 0))
        chosen = int(float(row.get("yixuan") or 0))
    except (TypeError, ValueError):
        return False, 0
    room = cap - chosen
    return room > 0, room


def capacity_text(row) -> str:
    cap = str(row.get("rongliang") or "?")
    chosen = str(row.get("yixuan") or "?")
    has_room, room = capacity_room(row)
    return f"余量 {cap} - 已选 {chosen}（{'有空位' if has_room else '满'}）"


def plan_preset_allocation(presets, occupied_times, catalog, targets=None, already_counts=None):
    """给每条预设挑一个教学班，使时间点分配合理。

    规则（顺序即志愿顺序）：
    1. 硬约束：任何候选都不能与**已选课程**时间冲突；
    2. 该类别还有"预计会选上"的名额时，还要避开**前面已经分配掉的志愿**时间，
       保证同一个时间段不会重复押注（通识最多 2 门这类上限也在这里体现）；
    3. 名额已被前面占满/本来就不缺的类别，只做"保底候选"——不占用时间段，
       但仍要求不与已选课程冲突；
    4. 候选之间优先选有余量的教学班，其次余量多的。

    返回 list[dict]，每项含 status：
      已分配          —— 可以与其他已分配志愿同时选上，时间不冲突
      备用            —— 保底用（类别名额已够，或与前序志愿同时间段）
      无可用时间      —— 所有候选都与已选课程冲突
      未找到          —— 本地缓存里没有这门课
    """
    targets = dict(targets or DEFAULT_TARGETS)
    already = dict(already_counts or {})
    # actual_room：现实还剩多少名额；planned：本次分配里已经"预定"给前序志愿的名额。
    # 两者分开，日志才能说清"是名额本来就没了"还是"被前序志愿占了"。
    actual_room = {category: max(0, int(targets.get(category, 0)) - int(already.get(category, 0)))
                   for category in targets}
    planned = {}
    occupied_times = [text for text in (occupied_times or []) if text]
    reserved = list(occupied_times)
    # 同一门课（课程号相同）的多个教学班只能押一个：教务那边同一课程选两个班，
    # 大概率是把先选上的那个替换掉（等于用一次机会换了个更差的班），也可能直接拒绝。
    # 所以第一个出现的志愿做主选，后面的同课程条目一律降为备用。
    first_of_course = {}
    results = []
    for position, raw in enumerate(presets):
        keyword, category = parse_preset(raw)
        entry = {"raw": raw, "keyword": keyword, "category": category, "status": "未找到",
                 "teaching_class": "", "time": "", "teacher": "", "capacity": "",
                 "reason": "", "candidates": []}
        matches = catalog_matches(catalog, keyword, category)
        if not matches:
            entry["reason"] = "本地「全部课程」缓存里没有这门课（先拉取一次全部课程）"
            results.append(entry)
            continue
        entry["candidates"] = [teaching_class for teaching_class, _row in matches]
        matches = sorted(matches, key=lambda item: (not capacity_room(item[1])[0],
                                                    -capacity_room(item[1])[1]))
        free = [(tc, row) for tc, row in matches
                if not any(time_conflicts(row.get("time"), text) for text in occupied_times)]
        if not free:
            teaching_class, row = matches[0]
            entry.update(teaching_class=teaching_class, time=clean_time_text(row.get("time")),
                         teacher=row.get("teacher", ""), capacity=capacity_text(row),
                         status="无可用时间", reason="所有候选都与已选课程时间冲突")
            results.append(entry)
            continue
        course_code = str(free[0][1].get("course_code") or free[0][1].get("kch") or "").strip().casefold()
        first_seen = first_of_course.get(course_code) if course_code else None
        if first_seen is not None:
            teaching_class, row = free[0]
            entry.update(teaching_class=teaching_class, time=clean_time_text(row.get("time")),
                         teacher=row.get("teacher", ""), capacity=capacity_text(row),
                         status="备用",
                         reason=f"与第 {first_seen + 1} 条同课程（{free[0][1].get('course_code') or course_code}），"
                                f"同一门课只押一个班，这条仅作保底（抢别的班）")
            results.append(entry)
            continue
        label = CATEGORY_PREFIX.get(category, "该类别")
        room_now = actual_room.get(category, 0) - planned.get(category, 0)
        picked = next(((tc, row) for tc, row in free
                       if not any(time_conflicts(row.get("time"), text) for text in reserved)), None)
        if room_now > 0 and picked:
            teaching_class, row = picked
            reserved.append(row.get("time"))
            planned[category] = planned.get(category, 0) + 1
            status, reason = "已分配", "与已选课程、前序志愿的时间段均不冲突"
        else:
            teaching_class, row = free[0]
            reasons = []
            if actual_room.get(category, 0) <= 0:
                reasons.append(f"{label}名额已够（{already.get(category, 0)}/{targets.get(category, '?')}）")
            elif room_now <= 0:
                reasons.append(f"{label}本次名额已被前序志愿占满")
            if picked is None:
                reasons.append("与前序志愿撞同一时间段")
            status = "备用"
            reason = "，".join(reasons + ["仅作保底"])
        if course_code:
            first_of_course.setdefault(course_code, position)
        entry.update(teaching_class=teaching_class, time=clean_time_text(row.get("time")),
                     teacher=row.get("teacher", ""), capacity=capacity_text(row),
                     status=status, reason=reason)
        results.append(entry)
    return results


def allocation_report(plan, occupied_courses=None) -> list:
    """把分配结果转成可直接写进日志的几行中文说明。"""
    lines = []
    if occupied_courses:
        busy = "；".join(f"{course.get('course_name') or course.get('teaching_class')}"
                         f"（{clean_time_text(course.get('time')) or '时间未知'}）"
                         for course in occupied_courses if course.get("time"))
        if busy:
            lines.append(f"[分配] 已选课程占用：{busy}")
    for entry in plan:
        head = f"[分配] {entry['status']}：{entry['raw']}"
        if entry.get("teaching_class"):
            head += f" → {entry['teaching_class']}"
        if entry.get("time"):
            head += f"｜{entry['time']}"
        if entry.get("capacity"):
            head += f"｜{entry['capacity']}"
        if entry.get("reason"):
            head += f"｜{entry['reason']}"
        lines.append(head)
    return lines


def clean_time_text(text) -> str:
    """教务时间/教室字段带 <br/> 和重复段，统一整理成"分号分隔、去重"的干净文本。"""
    value = re.sub(r"<br\s*/?>", "；", str(text or ""), flags=re.I)
    seen, kept = set(), []
    for part in re.split(r"[;；]", value):
        part = part.strip()
        if not part or part in seen:
            continue
        seen.add(part)
        kept.append(part)
    return "；".join(kept)


def parse_preset(raw: str):
    """Return (keyword, category) for a displayed preset line."""
    value = (raw or "").strip()
    if "|" in value:
        value = value.split("|", 1)[0].strip()
    display_tags = {"[通识]": "GE", "[英语]": "EN", "[体育]": "PE", "[课程]": "OTHER",
                    "通识": "GE", "英语": "EN", "体育": "PE", "课程": "OTHER"}
    for tag, category in display_tags.items():
        if value.startswith(tag):
            rest = value[len(tag):].strip()
            if "｜" in rest:
                fields = [field.strip() for field in rest.split("｜")]
                # 展示顺序：类别｜课程名｜教学班｜时间｜状态｜人数。
                return (fields[2] if len(fields) > 2 and fields[2] else
                        fields[1] if len(fields) > 1 else fields[0]), category
            return rest.split("|", 1)[0].strip(), category
    for prefix in ("GE", "EN", "PE"):
        if value.startswith(prefix + ":"):
            return value[len(prefix) + 1:].strip(), prefix
    return value, "OTHER"


def canonical_preset(raw: str) -> str:
    keyword, category = parse_preset(raw)
    return f"{category}:{keyword}" if category in {"GE", "EN", "PE"} else keyword


def match_pending_preset(pending: list[str], result_keyword: str, teaching_class: str) -> str:
    for wanted in (teaching_class, result_keyword):
        if wanted:
            match = next((raw for raw in pending if parse_preset(raw)[0] == wanted), None)
            if match:
                return match
    return teaching_class or result_keyword


def monitor_rush(base_url, cookie, presets, selector, interval, max_success, max_rounds, retry_rate, watch_full, parallel, auto_after_init, stop_event, log, on_success=None, initial_category_success=None, focus_count=0, full_wait=DEFAULT_FULL_WAIT, done_courses=None):
    client = make_client(base_url, cookie, log)
    presets = [p.strip() for p in presets if p.strip()]
    if not presets:
        raise RuntimeError("预设课程为空。")
    if parallel:
        parallel = False
        log("状态跟踪启用时自动使用顺序模式，避免重复课程并发提交。")
    log("监控启动：预设顺序 = 志愿顺序。")
    log(f"参数：interval={interval}s, max_success={max_success or '不限'}, max_rounds={max_rounds or '不限'}, "
        f"retry={retry_rate}, watch_full={watch_full}, focus_count={focus_count or '全部'}, full_wait={full_wait}s")
    success = 0
    category_success = {"GE": 0, "EN": 0, "PE": 0}
    category_success.update({k: int(v) for k, v in (initial_category_success or {}).items()
                             if k in category_success})
    # 教务已选 + 本程序已完成的**课程号**：同一门课的另一个教学班不要再抢。
    # 教务对同一课程选两个班通常是把先选上的替换掉（等于拿一次机会换了个更差的班），
    # 也可能直接拒绝——两种情况都是白费一次提交。
    done_courses = {str(code).strip() for code in (done_courses or []) if str(code).strip()}

    def course_code_of(course):
        """课程号（GENNET030 这种）。教务已选里也有同一个字段，两边口径一致。"""
        return str(getattr(course, "kch", "") or getattr(course, "course_code", "") or "").strip()

    if done_courses:
        log(f"教务已选/已完成的 {len(done_courses)} 门课程的课程号已计入去重（同一门课不会再抢第二个班）")
    category_limits = {"GE": 2, "EN": 1, "PE": 1}
    selected_courses, completed_raw = set(), set()
    latest_snapshot = {"completed": False, "progress": {}}
    round_no = 0
    # 盲提交缓存：keyword -> (course, target)。首次查询拿到 do_id/kch_id/xkkz_id 后，
    # 后续轮次直接复用提交，跳过查询（do_id 在轮次开放期间稳定）。
    # 服务器容量计数有刷新滞后，满课时不清缓存、持续盲打更能抢到退课空位。
    jxb_cache: dict = {}
    strikes: dict = {}   # 连续满课盲打次数，超过阈值强制重查一次刷新 do_id
    focus_count = max(0, int(focus_count or 0))
    try:
        full_wait = min(FULL_WAIT_MAX, max(FULL_WAIT_MIN, float(full_wait)))
    except (TypeError, ValueError):
        full_wait = DEFAULT_FULL_WAIT

    # ---- 收尾判定：一门志愿都提交不了时必须停车 ----
    # 实况（2026-09-17）：教务已选的通识已经 2 门（名额上限），而 18 个志愿全是通识 →
    # 每一轮每一门都被「名额已达到 2 门，跳过后续保底」跳过，一次提交都没发出去，
    # 循环却永远不停（用户：转了一晚上）。原来只有「抢到之后」才检查目标是否完成，
    # 「一个都抢不了」这种情况没有任何退出条件。
    # 现在：全量志愿里没有任何一门值得提交（名额满 / 已完成 / 与教务已选重复）→ 明确写日志退出。
    terminal_skips: set = set()      # 终局跳过的志愿（不是暂时失败，是再抢也没意义）
    submitted_this_round = 0         # 本轮真正发起过的提交次数

    def attemptable(raw_preset):
        """这门志愿还有提交价值吗？（只看已知状态，不联网、不碰缓存）"""
        if raw_preset in completed_raw:
            return False
        _keyword, _category = parse_preset(raw_preset)
        limit = category_limits.get(_category)
        if limit is not None and category_success.get(_category, 0) >= limit:
            return False
        return raw_preset not in terminal_skips

    def nothing_left_reason():
        """返回「没得抢了」的人话原因；还有可提交的志愿则返回 None。"""
        if any(attemptable(raw) for raw in presets):
            return None
        labels = {"GE": "通识", "EN": "英语", "PE": "体育"}
        bits = ["、".join(f"{labels.get(cat, cat)}名额已满（{category_success.get(cat, 0)}/{limit}）"
                          for cat, limit in category_limits.items()
                          if category_success.get(cat, 0) >= limit)]
        if completed_raw:
            bits.append(f"已完成 {len(completed_raw)} 门志愿")
        duplicate = len(terminal_skips - completed_raw)
        if duplicate:
            bits.append(f"{duplicate} 门与教务已选/已抢到的课程重复")
        detail = "；".join(bit for bit in bits if bit) or "所有志愿都已处理完"
        return f"没有可提交的志愿了（{detail}）—— 任务结束，抢课监控自动停止，不再空转。"

    log("类别名额基数（教务已选 + 本程序抢到）：" + "，".join(
        f"{ {'GE': '通识', 'EN': '英语', 'PE': '体育'}.get(cat, cat) } "
        f"{category_success.get(cat, 0)}/{limit}" for cat, limit in category_limits.items()))

    def finish_success(raw_preset, keyword, course, target):
        nonlocal success, latest_snapshot
        selected_courses.add(getattr(course, "kch_id", None))
        code = course_code_of(course)
        if code:
            done_courses.add(code)      # 同一门课的其它教学班随后都会被跳过
        completed_raw.add(raw_preset)
        success += 1
        if on_success:
            info = {
                "teaching_class": getattr(target, "jxbmc", "") or keyword,
                "time": getattr(target, "sksj", ""),
                "teacher": getattr(target, "jsxx", ""),
                "jxb_id": getattr(target, "jxb_id", ""),
                "kch_id": getattr(course, "kch_id", ""),
            }
            latest_snapshot = on_success(raw_preset, info) or latest_snapshot

    while not stop_event.is_set():
        round_no += 1
        if max_rounds and round_no > max_rounds:
            log("达到最大轮次，停止。")
            return latest_snapshot
        if max_success and success >= max_success:
            log("达到最多选 N 门，停止。")
            return latest_snapshot
        # 开抢前先看有没有得抢：没得抢就直接收工，一次网络请求都不发。
        stop_reason = nothing_left_reason()
        if stop_reason:
            log(stop_reason)
            return latest_snapshot
        log("")
        log(f"===== 第 {round_no} 轮 {datetime.now().strftime('%H:%M:%S')} =====")
        submitted_this_round = 0
        # 「优先盯前 N 门」：只高频盯最想要的几门，其余保底每 REST_EVERY_ROUNDS 轮试一次。
        # 前 N 门全部拿到（或该类名额已满）后自动恢复全量，让保底有机会顶上来。
        # 这是"把 70 多门课的时间集中在最想要的几门上"的办法——门数一多，串行的每一轮
        # 光排队就要几分钟，绝大多数时间都花在确认"那些课还是满的"。
        def focused_still_in_play():
            for raw in presets[:focus_count]:
                _kw, _cat = parse_preset(raw)
                if raw in completed_raw:
                    continue
                limit = category_limits.get(_cat)
                if limit is not None and category_success.get(_cat, 0) >= limit:
                    continue
                return True
            return False

        restrict = (focus_count > 0 and focus_count < len(presets)
                    and round_no % REST_EVERY_ROUNDS != 0 and focused_still_in_play())
        this_round = presets[:focus_count] if restrict else presets
        if restrict:
            log(f"[本轮范围] 只盯前 {focus_count} 门；" 
                f"其余 {len(presets) - focus_count} 门保底每 {REST_EVERY_ROUNDS} 轮试一次")
        elif focus_count > 0:
            log(f"[本轮范围] 全部 {len(presets)} 门（保底轮，或前 {focus_count} 门已到手）")
        if not init_until_open(client, log, once=False):
            time.sleep(interval)
            continue
        if not auto_after_init:
            log("init 已成功，但“init成功后自动抢课”未开启；监控停止在开放状态。")
            return latest_snapshot

        def attempt_one(raw_preset, use_cache=True):
            nonlocal submitted_this_round
            keyword = ""
            try:
                keyword, category = parse_preset(raw_preset)
                if raw_preset in completed_raw:
                    terminal_skips.add(raw_preset)
                    return raw_preset, keyword, False, "已完成，跳过", None
                limit = category_limits.get(category)
                if limit is not None and category_success[category] >= limit:
                    terminal_skips.add(raw_preset)
                    labels = {"GE": "通识", "EN": "英语", "PE": "体育"}
                    return raw_preset, keyword, False, f"{labels[category]}名额已达到{limit}门，跳过后续保底", None
                # ---- 盲提交路径：命中缓存则跳过查询，直接带 do_id 提交 ----
                cached = jxb_cache.get(keyword) if use_cache else None
                if cached:
                    course, target = cached
                    code = course_code_of(course)
                    if code and code in done_courses:
                        jxb_cache.pop(keyword, None)
                        strikes.pop(keyword, None)
                        terminal_skips.add(raw_preset)
                        return raw_preset, keyword, False, f"同课程已在教务/已完成（{code}），跳过另一个班", None
                    if getattr(course, "kch_id", None) in selected_courses:
                        jxb_cache.pop(keyword, None)
                        strikes.pop(keyword, None)
                        terminal_skips.add(raw_preset)
                        return raw_preset, keyword, False, "同课程已成功，跳过", None
                    submitted_this_round += 1
                    flag, msg = select_jxb(client, course, target, lambda m: None)
                    ok = str(flag) == "1"
                    if ok:
                        jxb_cache.pop(keyword, None)
                        strikes.pop(keyword, None)
                        if category in category_success:
                            category_success[category] += 1
                        return raw_preset, keyword, ok, msg, (course, target, flag, msg)
                    return raw_preset, keyword, False, annotate_select_msg(msg), None
                # ---- 查询路径：无缓存时取 kch_id/do_id/xkkz_id 并缓存 ----
                result = fetch_course(client, keyword, "auto")
                if not result.jxbs:
                    return raw_preset, keyword, False, "未找到教学班", None
                code = course_code_of(result.course)
                if code and code in done_courses:
                    terminal_skips.add(raw_preset)
                    return raw_preset, keyword, False, f"同课程已在教务/已完成（{code}），跳过另一个班", None
                if getattr(result.course, "kch_id", None) in selected_courses:
                    terminal_skips.add(raw_preset)
                    return raw_preset, keyword, False, "同课程已成功，跳过", None
                target = choose_jxb(result.jxbs, selector)
                if target is None:
                    return raw_preset, keyword, False, f"无匹配筛选：{selector}", None
                if use_cache:
                    jxb_cache[keyword] = (result.course, target)
                submitted_this_round += 1
                flag, msg = select_jxb(client, result.course, target, lambda m: None)
                ok = str(flag) == "1"
                if ok:
                    jxb_cache.pop(keyword, None)
                    strikes.pop(keyword, None)
                    if category in category_success:
                        category_success[category] += 1
                    return raw_preset, keyword, ok, msg, (result.course, target, flag, msg)
                return raw_preset, keyword, False, annotate_select_msg(msg), None
            except Exception as e:
                if keyword:
                    jxb_cache.pop(keyword, None)
                    strikes.pop(keyword, None)
                return raw_preset, keyword or str(raw_preset), False, f"{type(e).__name__}: {e}", None

        if parallel and len(this_round) > 1:
            with ThreadPoolExecutor(max_workers=min(4, len(this_round))) as pool:
                futures = [pool.submit(attempt_one, raw, False) for raw in this_round]
                for fut in as_completed(futures):
                    raw_preset, keyword, ok, msg, payload = fut.result()
                    log(f"{keyword}: {msg or ('成功' if ok else '失败')}")
                    if ok and payload:
                        course, target, _flag, _msg = payload
                        finish_success(raw_preset, keyword, course, target)
        else:
            for raw_preset in this_round:
                if stop_event.is_set():
                    return latest_snapshot
                raw_preset, keyword, ok, msg, payload = attempt_one(raw_preset)
                log(f"{keyword}: {msg or ('成功' if ok else '失败')}")
                if ok and payload:
                    course, target, _flag, _msg = payload
                    finish_success(raw_preset, keyword, course, target)
                    if latest_snapshot.get("completed"):
                        log("选课目标已全部完成。")
                        return latest_snapshot
                    continue
                cont, wait, reason = should_continue_for_msg(msg, retry_rate, watch_full)
                if cont:
                    if "容量已满" in reason and watch_full and keyword in jxb_cache:
                        # 满课蹲守：缓存不失效，节流盲打；服务器容量刷新滞后，
                        # 退课空位可能已存在但查询看不到，盲提交才能命中。
                        strikes[keyword] = strikes.get(keyword, 0) + 1
                        if strikes[keyword] >= 40:      # 盲打约 100 秒仍满：强制重查一次刷新 do_id
                            jxb_cache.pop(keyword, None)
                            strikes.pop(keyword, None)
                            log(f"{keyword}: 蹲守约 100 秒仍满，重新查询刷新参数")
                        else:
                            log(f"{keyword}: {reason}（盲打中，每 {full_wait}s 一次）")
                            time.sleep(full_wait)
                    elif wait:
                        log(reason)
                        time.sleep(wait)
                else:
                    # 硬失败（未找到/无匹配/其他）：清缓存，下一轮走查询刷新
                    jxb_cache.pop(keyword, None)
                    strikes.pop(keyword, None)
        if submitted_this_round == 0:
            # 本轮一次提交都没发出去：若原因是终局的（名额满/已完成/与教务已选重复），
            # 那就是真的没得抢了 —— 收工，别再空转一整晚。
            stop_reason = nothing_left_reason()
            if stop_reason:
                log(stop_reason)
                return latest_snapshot
        if max_success and success >= max_success:
            log("达到最多选 N 门，停止。")
            return latest_snapshot
        time.sleep(interval)
    log("用户停止。")
    return latest_snapshot


def export_courses_to_xlsx(rows, path: str):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "全部课程"
    headers = ["序号", "课程名", "课程号", "教学班名称", "学分", "教师", "上课时间", "上课教室", "已选/容量", "开课学院", "选课类别代码", "选课类别", "教学班ID", "课程ID"]
    ws.append(headers)
    for idx, row in enumerate(rows, 1):
        # 兼容当前 best-effort 行，以及未来完整课程行。
        cap = row.get("rongliang", "")
        yx = row.get("yixuan", "")
        sel = f"{yx}/{cap}" if (yx != "" or cap != "") else ""
        ws.append([
            idx,
            row.get("course_name", "") or row.get("keyword", ""),
            row.get("kch", ""),
            row.get("teaching_class", ""),
            row.get("xf", ""),
            row.get("teacher", ""),
            row.get("time", ""),
            row.get("room", ""),
            sel,
            row.get("xueyuan", ""),
            row.get("kklxdm", ""),
            row.get("course_tab", ""),
            row.get("jxb_id", ""),
            row.get("kch_id", ""),
        ])
    header_fill = PatternFill("solid", fgColor="10A37F")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = [8, 26, 12, 26, 8, 20, 34, 22, 12, 18, 36, 36]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    wb.save(path)


class App(tk.Tk):
    # 搜索抢课 Tab 筛选行的标签列宽（Tk 的 width 单位是字符宽，中文占 2 个槽）。
    # 10 槽 = 5 个汉字，和这一行原有的版式一致。超长的标签不要靠加宽列来解决：
    # 该行（课程归属）自己就有 7 个按钮，加宽标签会把右边的按钮挤出可见区
    # （实测默认 1120 窗口下「创新、创意、创业」「其它」会看不见）。
    # 用 "\n" 折成两行即可完整显示，且行宽不变。
    FILTER_LABEL_WIDTH = 10
    # 「整页滑动」容器（搜索抢课 Tab）的初始高度（px，会乘 DPI 缩放）。
    # 只是给 Tk 一个初值，真实高度由 grid 的 weight=1 撑满 Tab 剩余空间。
    SEARCH_PAGE_INIT_H = 200

    def __init__(self):
        super().__init__()
        self.title(f"{APP_NAME} {APP_VERSION}")
        try:
            self.iconbitmap(default=str(PROJECT_DIR / "zjgsu_launcher.ico"))
        except Exception:
            pass
        # Tk 已经根据 tk scaling 处理控件尺寸，窗口几何不要再次乘 DPI，
        # 否则高分屏会把底部日志和任务进度区域裁出屏幕。
        w = 1120
        h = 1000
        self.geometry(f"{w}x{h}")
        self.minsize(960, 620)
        # Tk 字体缩放：让文本按系统 DPI 渲染，避免高分屏发虚
        try:
            self.tk.call("tk", "scaling", 1.3333 * DPI_SCALE)
        except Exception:
            pass
        self.configure(bg=BG)
        self.q = queue.Queue()
        self.worker = None
        # 任务进度框是否已经收尾：False 表示有任务正在跑（表在转）。
        # poll() 用它做兜底，保证线程一结束状态不再停在“运行中”。
        self._progress_finished = True
        # 已选/进度实时对齐：空闲时每 auto_refresh_interval 秒对一次教务真相
        self.last_selected_refresh = None
        self.auto_refresh_interval = 60
        self.auto_refresh_var = None      # build 阶段创建
        self.avoid_conflict_var = None    # build 阶段创建
        self._poll_after_id = None
        self._tick_after_id = None
        self.stop_event = threading.Event()
        self.search_result = None
        self.all_courses_rows = []
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        LOG_DIR.mkdir(parents=True, exist_ok=True)
        self._migrate_data_file("course_presets.txt")
        self._migrate_data_file("selected_courses.json")
        self._migrate_data_file("system_selected_courses.json")
        self._migrate_data_file("course_catalog.json")
        self._migrate_data_file("all_courses_detailed.json")
        self._migrate_legacy_logs()
        self.log_file = LOG_DIR / "app.log"
        self.state_store = CourseStateStore(PRESET_FILE, SELECTED_STATE_FILE)
        self.course_catalog = self.load_course_catalog()
        self._catalog_core_index = None      # 课程库变了 → 核心教学班号索引重建
        self.system_selected = self.load_system_selected_cache()
        self.state_store.reconcile_system_selected(self.system_selected)
        # 自愈：已经抢到、却还留在待选列表里的条目（旧版本关键字写错导致没删掉）→ 清掉
        self._stale_completed = self.state_store.reconcile_completed_pending()
        self.setup_style()
        self.build()
        self.load_runtime_settings()
        if getattr(self, "_stale_completed", None):
            self.append("已自动清理待选列表里“已经抢到”的旧条目：" + "、".join(self._stale_completed))
        self._last_allocation_report = ""
        if self.avoid_conflict_var is not None and self.avoid_conflict_var.get():
            # 启动即按"已选课表 + 名额"对齐一次（没有变化时只写一行日志）
            self.allocate_presets(apply=True, reason="启动时对齐时间段")
        self.protocol("WM_DELETE_WINDOW", self.on_close)
        self.after(100, self.poll)
        self.after(5000, self.auto_refresh_tick)

    def setup_style(self):
        s = ttk.Style(self)
        try:
            s.theme_use("clam")
        except Exception:
            pass
        s.configure("TFrame", background=BG)
        s.configure("Card.TFrame", background=PANEL, bordercolor=SOFT, relief="solid")
        s.configure("TLabel", background=BG, foreground=TEXT, font=("Microsoft YaHei UI", 10))
        s.configure("Card.TLabel", background=PANEL, foreground=TEXT, font=("Microsoft YaHei UI", 10))
        s.configure("Muted.TLabel", background=BG, foreground=MUTED, font=("Microsoft YaHei UI", 9))
        s.configure("Title.TLabel", background=BG, foreground=TEXT, font=("Microsoft YaHei UI", 24, "bold"))
        s.configure("TButton", padding=(12, 7), background="#fff", foreground=TEXT, bordercolor=BORDER)
        s.configure("Accent.TButton", padding=(16, 9), background=ACCENT, foreground="#fff", bordercolor=ACCENT)
        s.configure("Danger.TButton", padding=(16, 9), background=DANGER, foreground="#fff", bordercolor=DANGER)
        s.configure("On.TButton", padding=(12, 7), background=ACCENT, foreground="#fff", bordercolor=ACCENT)
        s.configure("Off.TButton", padding=(12, 7), background="#ffffff", foreground=MUTED, bordercolor=BORDER)
        s.configure("Treeview", rowheight=int(30 * DPI_SCALE), background="#fff", foreground=TEXT, fieldbackground="#fff")
        s.configure("Treeview.Heading", background="#f1f5f9", foreground="#374151", font=("Microsoft YaHei UI", 10, "bold"))

    def build(self):
        root = ttk.Frame(self, padding=22)
        root.pack(fill="both", expand=True)
        root.columnconfigure(0, weight=1)
        root.rowconfigure(2, weight=1, minsize=260)
        root.rowconfigure(4, minsize=120)
        root.rowconfigure(6, minsize=38)
        ttk.Label(root, text=f"{APP_NAME} {APP_VERSION}", style="Title.TLabel").grid(row=0, column=0, sticky="w")
        ttk.Label(root, text="参考样板功能：Cookie 登录、课程代码精准查询、预设志愿、监控抢课、搜索抢课、全部课程尝试；界面选项已中文化。", style="Muted.TLabel").grid(row=1, column=0, sticky="w", pady=(4, 12))
        self.nb = ttk.Notebook(root)
        self.nb.grid(row=2, column=0, sticky="nsew")
        # 各 Tab 内部有较大的表格/课表请求高度；固定 Notebook 高度，
        # 给下方日志与任务进度区留出稳定、可见的空间。
        self.nb.configure(height=24)
        self.login_tab = ttk.Frame(self.nb, padding=14)
        self.auto_tab = ttk.Frame(self.nb, padding=14)
        self.search_tab = ttk.Frame(self.nb, padding=14)
        self.all_tab = ttk.Frame(self.nb, padding=14)
        self.selected_tab = ttk.Frame(self.nb, padding=14)
        self.nb.add(self.login_tab, text="登录")
        self.nb.add(self.auto_tab, text="自动抢课")
        self.nb.add(self.search_tab, text="搜索抢课")
        self.nb.add(self.all_tab, text="全部课程")
        self.nb.add(self.selected_tab, text="已选课程 / 进度")
        self.build_login()
        self.build_auto()
        self.build_search()
        self.build_all()
        self.build_selected()
        self.status = tk.StringVar(value="就绪")
        status_row = ttk.Frame(root)
        status_row.grid(row=3, column=0, sticky="ew", pady=(8, 4))
        ttk.Label(status_row, textvariable=self.status, style="Muted.TLabel").pack(side="left")
        ttk.Button(status_row, text="打开日志文件夹", command=self.open_log_folder).pack(side="right")
        self.log = tk.Text(root, height=9, bg="#fff", fg=TEXT, relief="solid", bd=1, highlightthickness=1, highlightbackground=SOFT, font=("Consolas", 10))
        self.log.grid(row=4, column=0, sticky="ew")
        ttk.Label(root, textvariable=tk.StringVar(value=f"运行日志：{self.log_file}"), style="Muted.TLabel").grid(row=5, column=0, sticky="w", pady=(3, 0))
        # 保持原始布局：日志下方是独立的任务进度框。
        progress_frame = tk.Frame(root, bg=SOFT, height=34, bd=1, relief="solid")
        progress_frame.grid(row=6, column=0, sticky="ew", pady=(6, 0))
        progress_frame.pack_propagate(False)
        self.task_progress_label = ttk.Label(progress_frame, text="任务进度：就绪", style="Muted.TLabel", width=18)
        self.task_progress_label.pack(side="left", padx=(8, 4))
        self.task_progress = ttk.Progressbar(progress_frame, mode="indeterminate", length=360)
        self.task_progress.pack(side="left", fill="x", expand=True, padx=(0, 8), pady=7)
        self.append("欢迎。建议：先登录并自动抓 Cookie，再到自动抢课页添加预设课程；可留空立即监控，也可填定时开始时间。")

    def build_login(self):
        f = self.login_tab
        ttk.Label(f, text="服务器地址", style="Card.TLabel").grid(row=0, column=0, sticky="w")
        self.base_var = tk.StringVar(value=BASE_URL_DEFAULT)
        ttk.Entry(f, textvariable=self.base_var, width=70).grid(row=0, column=1, sticky="ew", padx=10)
        ttk.Button(f, text="打开登录窗口", command=self.open_browser).grid(row=0, column=2, padx=4)
        ttk.Label(f, text="Cookie", style="Card.TLabel").grid(row=1, column=0, sticky="nw", pady=12)
        self.cookie_box = tk.Text(f, height=5, bg="#fff", fg=TEXT, relief="solid", bd=1, highlightthickness=1, highlightbackground=BORDER, font=("Consolas", 10))
        self.cookie_box.grid(row=1, column=1, columnspan=2, sticky="ew", padx=10, pady=12)
        btnrow = ttk.Frame(f)
        btnrow.grid(row=2, column=1, columnspan=2, sticky="w", padx=10)
        ttk.Button(btnrow, text="自动抓取 Cookie", style="Accent.TButton", command=self.capture_cookie).pack(side="left")
        ttk.Button(btnrow, text="复制 Cookie 给外部软件", command=self.copy_cookie).pack(side="left", padx=8)
        ttk.Button(btnrow, text="验证 Cookie", style="Accent.TButton", command=self.verify_cookie).pack(side="left")
        ttk.Button(btnrow, text="Cookie 登录 / 检查状态", command=self.check_cookie_login).pack(side="left", padx=8)
        ttk.Button(btnrow, text="关闭登录窗口 / 调试端口", command=self.close_browser).pack(side="left")
        f.columnconfigure(1, weight=1)

    def build_auto(self):
        f = self.auto_tab
        ttk.Label(f, text="预设输入", style="Card.TLabel").grid(row=0, column=0, sticky="w")
        self.preset_var = tk.StringVar()
        ttk.Entry(f, textvariable=self.preset_var, width=70).grid(row=0, column=1, columnspan=2, sticky="ew", padx=8)
        ttk.Button(f, text="添加", command=self.add_preset).grid(row=0, column=3)
        self.input_mode = tk.StringVar(value="自动识别")
        ttk.Combobox(f, textvariable=self.input_mode, values=["自动识别", "课程代码", "教学班", "课程名"], width=15, state="readonly").grid(row=0, column=4, padx=8)
        list_head = ttk.Frame(f)
        list_head.grid(row=1, column=0, columnspan=4, sticky="ew", pady=(10, 0))
        ttk.Label(list_head, text="待选课程", font=("Microsoft YaHei UI", 11, "bold")).pack(side="left")
        ttk.Label(list_head, text="类别前缀：GE=通识（最多2门）  EN=英语（最多1门）  PE=体育（最多1门）｜列表顺序=志愿顺序",
                  style="Muted.TLabel").pack(side="left", padx=16)
        ttk.Label(f, text="状态会在成功后自动更新；实时人数仅作参考，不作为淘汰条件。第三列＝抢课实际查询的关键字：写教学班号＝锁定该班；写课程名＝按名抢（搜索加入默认锁定）。",
                  style="Muted.TLabel", wraplength=980, justify="left").grid(row=2, column=0, columnspan=4, sticky="w")
        self.preset_list = tk.Listbox(f, height=10, bg="#fff", fg=TEXT, activestyle="dotbox", font=("Microsoft YaHei UI", 10))
        self.preset_list.grid(row=3, column=0, columnspan=2, sticky="nsew", pady=8)
        preset_scroll = ttk.Scrollbar(f, orient="vertical", command=self.preset_list.yview)
        preset_scroll.grid(row=3, column=2, sticky="nse", pady=8)
        self.preset_list.configure(yscrollcommand=preset_scroll.set)
        self.load_preset_details()
        self.load_presets_from_file()
        self.refresh_preset_display()
        btns = ttk.Frame(f)
        btns.grid(row=3, column=3, sticky="ns", pady=8)
        ttk.Button(btns, text="上移", command=lambda: self.move_preset(-1)).pack(fill="x")
        ttk.Button(btns, text="下移", command=lambda: self.move_preset(1)).pack(fill="x", pady=4)
        ttk.Button(btns, text="删除选中", command=self.delete_preset).pack(fill="x")
        ttk.Button(btns, text="清空全部", command=self.clear_presets).pack(fill="x", pady=4)
        ttk.Button(btns, text="刷新人数", command=self.refresh_live_counts).pack(fill="x", pady=4)
        self.interval_var = tk.StringVar(value="0.3")
        self.max_success_var = tk.StringVar(value="0")
        self.max_rounds_var = tk.StringVar(value="0")
        self.selector_var = tk.StringVar()
        row = ttk.Frame(f)
        row.grid(row=4, column=0, columnspan=4, sticky="ew", pady=8)
        for label, var, width in [("轮询间隔", self.interval_var, 7), ("最多选N门", self.max_success_var, 7), ("最大轮次", self.max_rounds_var, 7), ("时间/教师筛选", self.selector_var, 22)]:
            ttk.Label(row, text=label).pack(side="left", padx=(0, 4))
            ttk.Entry(row, textvariable=var, width=width).pack(side="left", padx=(0, 12))
        row2 = ttk.Frame(f)
        row2.grid(row=5, column=0, columnspan=3, sticky="ew", pady=(4, 0))
        self.start_time_var = tk.StringVar(value="")
        ttk.Label(row2, text="定时开始").pack(side="left", padx=(0, 4))
        ttk.Entry(row2, textvariable=self.start_time_var, width=12).pack(side="left")
        ttk.Label(row2, text="留空=立即监控；格式 09:59:58。未开放时会自动等开放，开放即抢。", foreground=MUTED).pack(side="left", padx=10)
        self.parallel_var = tk.BooleanVar(value=False)   # 旧版"并行抢课"开关，已废弃（见下）
        self.retry_var = tk.BooleanVar(value=True)
        self.watch_full_var = tk.BooleanVar(value=True)
        self.auto_after_init_var = tk.BooleanVar(value=True)
        # 长期机制开关：加入课程 / 开抢前自动按已选课表避开时间段冲突
        self.avoid_conflict_var = tk.BooleanVar(value=True)
        opts = ttk.Frame(f)
        opts.grid(row=6, column=0, columnspan=3, sticky="w", pady=(8, 0))
        self.retry_btn = ttk.Button(opts, command=lambda: self.toggle_option(self.retry_var, self.retry_btn, "频率过高自动重试"))
        self.retry_btn.pack(side="left")
        self.watch_full_btn = ttk.Button(opts, command=lambda: self.toggle_option(self.watch_full_var, self.watch_full_btn, "满课蹲守退课"))
        self.watch_full_btn.pack(side="left", padx=12)
        self.auto_after_init_btn = ttk.Button(opts, command=lambda: self.toggle_option(self.auto_after_init_var, self.auto_after_init_btn, "init成功后自动抢课"))
        self.auto_after_init_btn.pack(side="left", padx=12)
        self.refresh_toggle(self.retry_var, self.retry_btn, "频率过高自动重试")
        self.refresh_toggle(self.watch_full_var, self.watch_full_btn, "满课蹲守退课")
        self.refresh_toggle(self.auto_after_init_var, self.auto_after_init_btn, "init成功后自动抢课")
        # 旧「并行抢课」按钮已删除：monitor_rush 里它本来就被无条件关掉（并发提交会让
        # "哪一门成功了"记不清账），留着只会骗人。换成两个真正有用的旋钮：
        #   优先盯前 N 门 —— 门数多时只高频盯最想要的几门（其余每 3 轮试一次）
        #   满课节流     —— 满课盲打的间隔，2.5 秒是保守值，调小更快但可能被限流
        self.focus_count_var = tk.StringVar(value="0")
        self.full_wait_var = tk.StringVar(value=f"{DEFAULT_FULL_WAIT:g}")
        # 专门开一行：挤在上一行会把「自动抢课」Tab 的请求宽度从 1381 顶到 1875（高分屏下被裁切）
        tune_row = ttk.Frame(f)
        tune_row.grid(row=8, column=0, columnspan=4, sticky="w", pady=(6, 0))
        ttk.Label(tune_row, text="优先盯前N门").pack(side="left", padx=(0, 4))
        ttk.Entry(tune_row, textvariable=self.focus_count_var, width=4).pack(side="left")
        ttk.Label(tune_row, text="满课节流(秒)").pack(side="left", padx=(16, 4))
        ttk.Entry(tune_row, textvariable=self.full_wait_var, width=5).pack(side="left")
        ttk.Label(tune_row, text="0=盯全部；门数多时只高频盯最想要的几门，其余保底每 3 轮试一次。"
                                  "满课节流 2.5 秒是保守值，调小更激进但可能被教务判「频率过高」。",
                  style="Muted.TLabel", wraplength=560, justify="left").pack(side="left", padx=10)
        # 时间段分配单独一行：避免把本 Tab 撑得过宽（高分屏下右侧控件会被裁切）
        row3 = ttk.Frame(f)
        row3.grid(row=7, column=0, columnspan=4, sticky="w", pady=(6, 0))
        self.avoid_conflict_btn = ttk.Button(row3, command=lambda: self.toggle_option(self.avoid_conflict_var, self.avoid_conflict_btn, "自动避开时间冲突"))
        self.avoid_conflict_btn.pack(side="left")
        self.refresh_toggle(self.avoid_conflict_var, self.avoid_conflict_btn, "自动避开时间冲突")
        ttk.Button(row3, text="自动分配时间段（按已选课表）",
                   command=lambda: self.allocate_presets(apply=True, reason="手动点击")).pack(side="left", padx=10)
        ttk.Label(row3, text="按「已选课程时间 + 各类别剩余名额」给每条志愿挑教学班；加入课程和开抢前也会自动跑一次。",
                  style="Muted.TLabel", wraplength=460, justify="left").pack(side="left", padx=10)
        ttk.Button(f, text="开始监控", style="Danger.TButton", command=self.start_monitor).grid(row=9, column=0, pady=12, sticky="w")
        ttk.Button(f, text="停止", command=self.stop_monitor).grid(row=9, column=1, pady=12, sticky="w")
        f.columnconfigure(1, weight=1)
        f.rowconfigure(3, weight=1)

    def build_search(self):
        f = self.search_tab
        self.search_var = tk.StringVar()
        self.search_mode = tk.StringVar(value="自动识别")
        self.search_loop_var = tk.BooleanVar(value=False)
        self._refresh_availability_next = False
        self.search_filters = {"course_categories": [], "course_belongs": [], "course_groups": [], "days": [], "periods": [],
                               "availability": "全部", "campus": CAMPUS_LOCAL,
                               "weeks": "", "weeks_mode": WEEK_MODE_INCLUDE, "credits": []}
        self.search_filter_buttons = []
        self._search_row_index = {}     # Treeview iid → 在搜索结果里的序号（列顺序变了也不会错位）
        self._search_checked = set()    # 已勾选的行 iid（复选框）
        ttk.Label(f, text="课程名/课程号/教学班").grid(row=0, column=0, sticky="w")
        search_entry = ttk.Entry(f, textvariable=self.search_var, width=44)
        search_entry.grid(row=0, column=1, sticky="ew", padx=8)
        search_entry.bind("<FocusOut>", lambda _event: self.save_search_settings())
        search_mode_box = ttk.Combobox(f, textvariable=self.search_mode, values=["自动识别", "课程代码", "教学班", "课程名"], width=15, state="readonly")
        search_mode_box.grid(row=0, column=2)
        search_mode_box.bind("<<ComboboxSelected>>", lambda _event: self.save_search_settings())
        ttk.Button(f, text="搜索", style="Accent.TButton", command=self.search_course).grid(row=0, column=3, padx=6)
        ttk.Button(f, text="一键抢课（勾选行）", style="Danger.TButton", command=self.submit_search_selected).grid(row=0, column=4, padx=6)
        self.add_preset_btn = ttk.Button(f, text="加入勾选到抢课", command=self.add_search_to_preset)
        self.add_preset_btn.grid(row=0, column=5, padx=6)
        # ---- 整页可滑动（v0.1.8，用户要求：让整个「搜索抢课」页面整体滑动）--------
        # 背景：这一页的内容比 Notebook 高得多 —— 实测 1920x1040 下结果表格只剩 1px 高。
        # 做法：搜索行留在顶部不动，其余（筛选区 + 循环抢课 + 勾选工具行 + 结果表格）
        # 整块放进一个可滚动页面：页面不够高就整体上下滑，表格拿回它该有的高度，
        # 也不会出现"筛选区里再套一层滚动"的双层滚动。
        page_host = ttk.Frame(f)
        page_host.grid(row=1, column=0, columnspan=6, sticky="nsew", pady=(8, 0))
        f.rowconfigure(1, weight=1)
        self.page_canvas = tk.Canvas(page_host, highlightthickness=0, bd=0, bg=BG,
                                     height=int(self.SEARCH_PAGE_INIT_H * DPI_SCALE))
        self.page_scroll = ttk.Scrollbar(page_host, orient="vertical", command=self.page_canvas.yview)
        self.page_canvas.configure(yscrollcommand=self.page_scroll.set)
        self.page_canvas.pack(side="left", fill="both", expand=True)
        self.page_scroll.pack(side="right", fill="y")
        page = ttk.Frame(self.page_canvas)
        self._page_frame = page
        self._page_window = self.page_canvas.create_window((0, 0), window=page, anchor="nw")
        page.bind("<Configure>", lambda _event: self._sync_page_scroll())
        self.page_canvas.bind("<Configure>", self._on_page_canvas_configure)
        # Windows 上 <MouseWheel> 只送给焦点控件 → 全局收事件，再按指针位置决定滚谁：
        # 指针在结果表格上（且表格自己还有内容可滚）就滚表格，否则滚整页；
        # 指针不在本页（待选列表 / 课表）时原样放行。
        self.bind_all("<MouseWheel>", self._on_page_wheel, add="+")
        filter_box = ttk.LabelFrame(page, text="筛选（点击标签，可多选）", padding=6)
        filter_box.pack(fill="x")
        self._filter_box = filter_box
        # 「更多筛选」折叠：课程归属 + 课程组这两行最占地方，收进折叠区，默认收起。
        # 里面有启用的条件时按钮上会写「已启用 N 项」，不会悄悄过滤掉用户的课程。
        self.filter_more_open = tk.BooleanVar(value=False)
        self._more_row = ttk.Frame(filter_box)
        self._more_row.pack(fill="x", pady=(0, 2))
        self.more_filter_btn = ttk.Button(self._more_row, text="▸ 更多筛选（课程归属 / 课程组）",
                                          command=self.toggle_more_filters)
        self.more_filter_btn.pack(side="left")
        ttk.Label(self._more_row, text="通识细分与课程分组在这里面；点开勾选，收起后表格腾出高度",
                  style="Muted.TLabel").pack(side="left", padx=8)
        # 标签「课程归属\n（通识）」折成两行：Tk 的 Label 不会自动折行，写成一行的
        # 话超宽部分会被裁成「课程归属（通」；而把标签列加宽（10 槽 → 12 槽以上）又会
        # 把这一行右边的按钮挤出可见区（这一行本来就有 7 个按钮）。折行既不裁字也不占额外宽度。
        belongs_row = self._build_filter_tags(filter_box, "课程归属\n（通识）", ["文学、历史、哲学", "艺术、宗教、文化", "经济、管理、法律", "写作、认知、表达", "自然、工程、技术", "创新、创意、创业", "其它"], "course_belongs")
        group_box = ttk.LabelFrame(filter_box, text="课程组（按类别分组）", padding=3)
        group_box.pack(fill="x", pady=2)
        group_row = ttk.Frame(group_box)
        group_row.pack(fill="x", pady=2)
        for label in COURSE_GROUP_ALIASES:
            button = tk.Button(group_row, text=label, relief="flat", bd=0, padx=7, pady=2,
                               bg="#f1f5f9", fg=TEXT, activebackground="#d1fae5",
                               command=lambda v=label: self.toggle_search_filter("course_groups", v))
            button.pack(side="left", padx=2)
            self.search_filter_buttons.append(("course_groups", label, button))
        # 折叠区就这两行（最占地方的两行）：收起时表格立刻多出约 100px。
        self._more_rows = (belongs_row, group_box)
        self._build_filter_tags(filter_box, "上课星期", ["一", "二", "三", "四", "五", "六", "日"], "days")
        self._build_filter_tags(filter_box, "上课节次", [str(i) for i in range(1, 15)], "periods")
        # 周数筛选：教务的课常只上其中几周（{8-13周}），同一时间段但周次不重叠的两门课
        # 不会真的撞车。输入 1-7 →「只留含于」留下 1-3、2-7（1-16 这种超集被去掉）；
        # 点按钮切到「剔除」后输入 8-13 → 只留 1-7、14-16。
        weeks_row = ttk.Frame(filter_box)
        weeks_row.pack(fill="x", pady=2)
        ttk.Label(weeks_row, text="周数筛选", width=self.FILTER_LABEL_WIDTH).pack(side="left")
        self.search_weeks_var = tk.StringVar()
        weeks_entry = ttk.Entry(weeks_row, textvariable=self.search_weeks_var, width=14)
        weeks_entry.pack(side="left", padx=(0, 4))
        weeks_entry.bind("<Return>", lambda _event: self.apply_search_weeks())
        weeks_entry.bind("<FocusOut>", lambda _event: self.apply_search_weeks())
        self.search_weeks_mode_btn = ttk.Button(weeks_row, text="模式：只留含于",
                                                command=self.toggle_search_weeks_mode)
        self.search_weeks_mode_btn.pack(side="left", padx=4)
        ttk.Button(weeks_row, text="清空", command=self.clear_search_weeks).pack(side="left", padx=4)
        self.search_weeks_hint = tk.StringVar()
        ttk.Label(weeks_row, textvariable=self.search_weeks_hint, style="Muted.TLabel").pack(side="left", padx=8)
        # 学分筛选：学分是教务返回行里的 xf。按钮按课程数据里**实际出现过的学分**动态生成
        # （不会出现"点了永远搜不到"的死按钮）；通识选修课只有 1 分 / 2 分两档。
        credits_row = ttk.Frame(filter_box)
        credits_row.pack(fill="x", pady=2)
        ttk.Label(credits_row, text="学分筛选", width=self.FILTER_LABEL_WIDTH).pack(side="left")
        self.credit_button_box = ttk.Frame(credits_row)
        self.credit_button_box.pack(side="left")
        ttk.Button(credits_row, text="清空", command=self.clear_search_credits).pack(side="left", padx=4)
        self.credit_buttons = []
        self.rebuild_credit_buttons()
        self._build_filter_tags(filter_box, "有无余量", ["全部", "有余量", "无余量"], "availability", single=True)
        # 校区：下沙=我的校区。教工路校区的课单程 40 分钟以上，默认直接过滤掉。
        self._build_filter_tags(filter_box, "上课校区", [CAMPUS_ALL, CAMPUS_LOCAL, "教工路"], "campus", single=True)
        self.search_loop_btn = ttk.Button(page, command=lambda: self.toggle_option(self.search_loop_var, self.search_loop_btn, "循环抢课（每1.5秒重试）"))
        self.search_loop_btn.pack(fill="x", pady=6)
        self.refresh_toggle(self.search_loop_var, self.search_loop_btn, "循环抢课（每1.5秒重试）")
        # 批量勾选工具行（单独一行：挤在同一行会把 Tab 撑宽、高分屏下右侧控件被裁切）
        pick_row = ttk.Frame(page)
        pick_row.pack(fill="x", pady=(0, 4))
        ttk.Button(pick_row, text="☑ 全选", command=lambda: self.check_all_search(True)).pack(side="left")
        ttk.Button(pick_row, text="☐ 清空勾选", command=lambda: self.check_all_search(False)).pack(side="left", padx=6)
        ttk.Button(pick_row, text="⇄ 反选", command=self.invert_search_checks).pack(side="left")
        self.pick_hint = ttk.Label(pick_row, text="点第一列方框可勾选多行 → 批量加入抢课 / 批量抢课",
                                   style="Muted.TLabel", wraplength=620)
        self.pick_hint.pack(side="left", padx=12)
        cols = ("pick", "idx", "name", "xf", "class", "campus", "time", "teacher", "status")
        self.search_tree = ttk.Treeview(page, columns=cols, show="headings", height=10, selectmode="extended")
        # 列宽：新增「学分」列后，其余列各让出一点，避免整个 Tab 的请求宽度继续变胖。
        # 「选」列收窄到 26px（用户要求：列窄一点），但勾选行整行绿色加粗（checked tag），
        # 让勾没勾一眼可辨——Tk 的 Treeview 改不了单列字号，整行高亮是唯一能真正"更醒目"的手段。
        for c, t, w in [("pick", "选", 26), ("idx", "序号", 38), ("name", "课程名", 180),
                        ("xf", "学分", 46), ("class", "教学班名称", 216), ("campus", "校区", 70),
                        ("time", "上课时间", 252), ("teacher", "教师", 140), ("status", "操作提示", 88)]:
            self.search_tree.heading(c, text=t)
            self.search_tree.column(c, width=w, anchor="center" if c in ("pick", "idx", "xf", "campus", "status") else "w")
        self.search_tree.tag_configure("offcampus", foreground="#b91c1c")
        self.search_tree.tag_configure("checked", foreground=ACCENT, font=("Microsoft YaHei UI", 9, "bold"))
        # 表头「选」也可以点：全选/全不选（整表切换，不用一行行点）
        self.search_tree.heading("pick", text="选", command=self.toggle_select_all)
        self.search_tree.pack(fill="x", pady=(10, 0))
        self.search_tree.bind("<Button-1>", self._on_search_tree_click)
        self.search_tree.bind("<space>", lambda _event: self.toggle_search_check())
        self.search_tree.bind("<Control-a>", self._on_search_select_all_key)
        f.columnconfigure(1, weight=1)
        self.load_search_settings()
        self._apply_more_state()
        self.after_idle(self._sync_page_scroll)

    def _build_filter_tags(self, parent, label, values, key, single=False):
        row = ttk.Frame(parent)
        row.pack(fill="x", pady=2)
        # 标签列宽统一（10 个字符槽，中文占 2 槽）；标签文字里可以用 "\n" 折行，
        # 别把列加宽——这一行右边还有一排按钮，列一宽按钮就被挤出可见区。
        ttk.Label(row, text=label, width=self.FILTER_LABEL_WIDTH).pack(side="left")
        for value in values:
            shown = (f"星期{value}" if key == "days" else value)
            button = tk.Button(row, text=shown, relief="flat", bd=0, padx=7, pady=2,
                               bg="#f1f5f9", fg=TEXT, activebackground="#d1fae5",
                               command=lambda v=value, k=key, s=single: self.toggle_search_filter(k, v, s))
            button.pack(side="left", padx=2)
            self.search_filter_buttons.append((key, value, button))
        return row     # 折叠/滚动要能拿到整行控件

    # --- 筛选区：可滚动 + 「更多筛选」折叠 -----------------------------------
    def _folded_condition_count(self):
        """折叠区里已启用的条件数（课程归属 + 课程组）。"""
        return sum(len(self.search_filters.get(key) or []) for key in ("course_belongs", "course_groups"))

    def _update_more_label(self):
        """折叠按钮的文字始终如实写明状态：收起时也不掩盖"里面有条件在生效"。"""
        open_ = bool(self.filter_more_open.get())
        text = "▾ 更多筛选（课程归属 / 课程组）" if open_ else "▸ 更多筛选（课程归属 / 课程组）"
        count = self._folded_condition_count()
        if count:
            text += f" · 已启用 {count} 项"
        try:
            self.more_filter_btn.configure(text=text)
        except (AttributeError, tk.TclError):
            pass

    def toggle_more_filters(self):
        """点「更多筛选」：收起/摊开课程归属 + 课程组（状态会记住，下次打开还在）。"""
        self.filter_more_open.set(not bool(self.filter_more_open.get()))
        self._apply_more_state()
        self.save_search_settings()

    def _apply_more_state(self, *_args):
        """按 filter_more_open 显示/隐藏折叠行，并重算筛选区高度。"""
        self._update_more_label()
        rows = getattr(self, "_more_rows", ())
        for widget in rows:
            try:
                widget.pack_forget()
            except tk.TclError:
                pass
        if bool(self.filter_more_open.get()):
            anchor = self._more_row
            for widget in rows:
                # after= 保证摊开后仍在「更多筛选」那一行下面，顺序不乱。
                widget.pack(fill="x", pady=2, after=anchor)
                anchor = widget
        self.after_idle(self._sync_page_scroll)

    def _on_page_canvas_configure(self, event):
        """画布变宽 → 里面的整页跟着变宽（不然内容会按请求宽度显示、右边被裁）。"""
        try:
            self.page_canvas.itemconfigure(self._page_window, width=event.width)
        except tk.TclError:
            return
        self._sync_page_scroll()

    def _page_scroll_visible(self):
        """整页滚动条是不是摆出来了。

        winfo_ismapped 在窗口隐藏时永远是 False → 用 winfo_manager 判断"有没有被 pack"，
        与窗口可见性无关，测试也好断言。
        """
        try:
            return self.page_scroll.winfo_manager() == "pack"
        except (AttributeError, tk.TclError):
            return False

    def _sync_page_scroll(self):
        """刷新整页滚动区域：装得下就把滚动条收起来，装不下就摆出来自己滚。"""
        canvas = getattr(self, "page_canvas", None)
        if canvas is None or not canvas.winfo_exists():
            return
        page = getattr(self, "_page_frame", None)
        content = page.winfo_reqheight() if page is not None and page.winfo_exists() else 0
        view = canvas.winfo_height()
        canvas.configure(scrollregion=(0, 0, canvas.winfo_width(), max(content, view)))
        need = content > view + 1
        if need and not self._page_scroll_visible():
            self.page_scroll.pack(side="right", fill="y")
        elif not need and self._page_scroll_visible():
            self.page_scroll.pack_forget()
        if not need:
            canvas.yview_moveto(0)

    def _on_page_wheel(self, event):
        """滚轮：指针在结果表格上、且表格自己还有内容可滚 → 滚表格；否则滚整页。

        指针不在本页（待选列表 / 课表）时原样放行，各滚各的。
        """
        canvas = getattr(self, "page_canvas", None)
        if canvas is None or not canvas.winfo_exists():
            return None
        try:
            x, y = (event.x_root, event.y_root) if event is not None else self.winfo_pointerxy()
            node = canvas.winfo_containing(x, y)
        except tk.TclError:
            return None
        walker = node
        while walker is not None:
            if walker is canvas:
                break
            walker = getattr(walker, "master", None)
        else:
            return None                       # 指针不在本页
        if node is self.search_tree:
            first, last = self.search_tree.yview()
            if last - first < 1.0 - 1e-9:     # 表格自己还能滚 → 让表格滚
                return None
        canvas.yview_scroll(int(-event.delta / 120), "units")
        return "break"

    @staticmethod
    def _migrate_group_value(value):
        """把旧版存储的细分课程组值（如"体育分项""军事理论"）归并到新组名。"""
        for label, aliases in COURSE_GROUP_ALIASES.items():
            if value == label or value in aliases:
                return label
        return value

    # 单选型筛选：这些键存的是单个字符串（不是列表）。
    SINGLE_VALUE_FILTERS = ("availability", "campus")

    def toggle_search_filter(self, key, value, single=False):
        if key in self.SINGLE_VALUE_FILTERS:
            self.search_filters[key] = value
        elif single:
            self.search_filters[key] = [value]
        else:
            selected = self.search_filters.setdefault(key, [])
            if value in selected:
                selected.remove(value)
            else:
                selected.append(value)
        self.refresh_search_filter_buttons()
        self.save_search_settings()
        if key == "availability":
            # 余量是易变数据：每次点击该组标签都触发在线刷新，失败再回退缓存。
            self._refresh_availability_next = True
            self.search_course()
        elif key == "campus":
            # 校区是本地就能判定的：点一下立刻按新校区重跑本地搜索，不用等联网。
            self.search_course()

    def refresh_search_filter_buttons(self):
        for key, value, button in list(self.search_filter_buttons):
            try:
                if key in self.SINGLE_VALUE_FILTERS:
                    selected = value == self.search_filters.get(key)
                else:
                    selected = value in self.search_filters.get(key, [])
                button.configure(bg="#10a37f" if selected else "#f1f5f9", fg="#ffffff" if selected else TEXT)
            except tk.TclError:
                # 学分行重建时旧按钮已销毁：顺手摘掉，别让列表里留死引用。
                self.search_filter_buttons = [item for item in self.search_filter_buttons if item[2] is not button]
        # 折叠区里的条件变了 → 按钮上的「已启用 N 项」要跟着变（收起时也不骗人）。
        self._update_more_label()

    # --- 学分筛选（按数据里真实出现过的学分动态生成按钮） -------------------
    def available_credits(self) -> list:
        """当前课程数据里出现过的学分（升序）。都没有就退化成通识课最常见两档。"""
        seen = set()
        for row in (getattr(self, "all_courses_rows", None) or []):
            value = credit_value(row)
            if value:
                seen.add(value)
        if not seen:
            for row in (getattr(self, "course_catalog", None) or {}).values():
                value = credit_value(row)
                if value:
                    seen.add(value)
        if not seen:
            return ["1", "2"]

        def sort_key(value):
            try:
                return (0, float(value))
            except ValueError:
                return (1, 0.0)

        return sorted(seen, key=sort_key)

    def rebuild_credit_buttons(self):
        """重建学分行按钮：本地缓存或在线刷新让学分档位变化后，不留会骗人的旧按钮。"""
        if not hasattr(self, "credit_button_box"):
            return
        for button in self.credit_buttons:
            button.destroy()
        self.search_filter_buttons = [item for item in self.search_filter_buttons if item[0] != "credits"]
        self.credit_buttons = []
        for value in self.available_credits():
            button = tk.Button(self.credit_button_box, text=f"{value}分", relief="flat", bd=0,
                               padx=7, pady=2, bg="#f1f5f9", fg=TEXT, activebackground="#d1fae5",
                               command=lambda v=value: self.toggle_search_filter("credits", v))
            button.pack(side="left", padx=2)
            self.search_filter_buttons.append(("credits", value, button))
            self.credit_buttons.append(button)
        # 旧设置里可能存着当前数据中不存在的学分档位——那会变成"看不见却生效"的过滤，
        # 直接剔掉（界面上虽有「清空」，但用户看不出是哪个档位在起作用）。
        valid = set(self.available_credits())
        stale = [v for v in (self.search_filters.get("credits") or []) if v not in valid]
        if stale:
            self.search_filters["credits"] = [v for v in self.search_filters["credits"] if v in valid]
            try:
                self.append(f"学分筛选里的「{'/'.join(stale)} 分」在当前课程数据里不存在，已自动清掉。")
            except Exception:
                pass
        self.refresh_search_filter_buttons()

    def clear_search_credits(self):
        """清空学分筛选并按新条件重跑本地搜索（和「上课校区」一样不用联网）。"""
        if not self.search_filters.get("credits"):
            return
        self.search_filters["credits"] = []
        self.refresh_search_filter_buttons()
        self.save_search_settings()
        self.search_course()

    # --- 周数筛选（「周数筛选」输入框 + 模式按钮） ---------------------------
    def refresh_search_weeks_controls(self):
        """刷新「周数筛选」那一行：模式文字 + 一行说人话的效果说明。

        不切换按钮 style：Accent.TButton 的 padding 更大，来回切换会让行高跳动。
        """
        mode = self.search_filters.get("weeks_mode") or WEEK_MODE_INCLUDE
        text = str(self.search_filters.get("weeks") or "").strip()
        self.search_weeks_mode_btn.configure(
            text="模式：" + WEEK_MODE_BUTTON_TEXT.get(mode, WEEK_MODE_BUTTON_TEXT[WEEK_MODE_INCLUDE]))
        if not text:
            self.search_weeks_hint.set("1-7、8-13，回车生效")
        elif mode == WEEK_MODE_EXCLUDE:
            self.search_weeks_hint.set(f"已生效：去掉与 {text} 周重叠的课")
        else:
            self.search_weeks_hint.set(f"已生效：只留周次在 {text} 内的课")

    def toggle_search_weeks_mode(self):
        """在「只留含于」和「剔除重叠」之间切换。"""
        current = self.search_filters.get("weeks_mode") or WEEK_MODE_INCLUDE
        self.search_filters["weeks_mode"] = (WEEK_MODE_EXCLUDE if current == WEEK_MODE_INCLUDE
                                             else WEEK_MODE_INCLUDE)
        self.save_search_settings()
        self.refresh_search_weeks_controls()
        self.status.set(f"周数筛选模式：{WEEK_MODE_BUTTON_TEXT[self.search_filters['weeks_mode']]}"
                        "（输入周数后回车生效）")
        self.search_course()

    def clear_search_weeks(self):
        self.search_weeks_var.set("")
        self.apply_search_weeks()

    def apply_search_weeks(self):
        """把输入框里的周数条件写进筛选设置并重跑本地搜索（和「上课校区」一样，不用联网）。"""
        text = self.search_weeks_var.get().strip()
        if text and not week_set(text):
            # 解析不出就清空：留着条件却不生效，比没有更让人困惑。
            self.append(f"周数「{text}」识别不出周次（示例：1-7、8-13、1-7,14-16，可用单双周），已清空该条件。")
            self.status.set("周数条件无法识别（示例：1-7、8-13），已清空")
            self.search_weeks_var.set("")
            text = ""
        changed = text != str(self.search_filters.get("weeks") or "")
        self.search_filters["weeks"] = text
        self.save_search_settings()
        self.refresh_search_weeks_controls()
        if changed:
            self.search_course()

    def load_search_settings(self):
        try:
            data = json.loads(SETTINGS_FILE.read_text(encoding="utf-8")) if SETTINGS_FILE.exists() else {}
            saved = data.get("search_filters", {})
            for key in ("course_categories", "course_belongs", "course_groups", "days", "periods", "credits"):
                if key == "course_categories":
                    # UI 已移除「课程类别」行（与课程组重复）；旧设置里的值不可见也无法取消，
                    # 读入会造成隐式过滤，因此一律丢弃。
                    self.search_filters[key] = []
                    continue
                saved_values = list(saved.get(key, []))
                if key == "course_groups":
                    saved_values = list(dict.fromkeys(self._migrate_group_value(v) for v in saved_values))
                self.search_filters[key] = saved_values
            self.search_filters["availability"] = saved.get("availability", "全部")
            campus = saved.get("campus", CAMPUS_LOCAL)
            if isinstance(campus, (list, tuple)):     # 兼容早期误存成列表的设置
                campus = campus[0] if campus else CAMPUS_LOCAL
            self.search_filters["campus"] = campus or CAMPUS_LOCAL
            self.search_filters["weeks"] = str(saved.get("weeks") or "")
            saved_weeks_mode = saved.get("weeks_mode")
            self.search_filters["weeks_mode"] = (saved_weeks_mode
                                                 if saved_weeks_mode in (WEEK_MODE_INCLUDE, WEEK_MODE_EXCLUDE)
                                                 else WEEK_MODE_INCLUDE)
            self.search_weeks_var.set(self.search_filters["weeks"])
            self.search_var.set(data.get("search_keyword", ""))
            self.search_mode.set(data.get("search_mode", "自动识别"))
            # 「更多筛选」折叠状态：记住用户上次的选择；但折叠区里有启用中的条件时自动摊开，
            # 免得"看不见的筛选条件"把课程悄悄过滤掉（用户的 course_groups 就是常开的）。
            saved_more = bool(data.get("filter_more_open", False))
            self.filter_more_open.set(saved_more or self._folded_condition_count() > 0)
            self._apply_more_state()
            self.refresh_search_filter_buttons()
            self.refresh_search_weeks_controls()
        except Exception as e:
            self.append(f"读取搜索设置失败：{type(e).__name__}: {e}")

    def save_search_settings(self):
        try:
            data = {}
            if SETTINGS_FILE.exists():
                data = json.loads(SETTINGS_FILE.read_text(encoding="utf-8"))
            data["search_filters"] = self.search_filters
            data["search_keyword"] = self.search_var.get()
            data["search_mode"] = self.search_mode.get()
            data["filter_more_open"] = bool(self.filter_more_open.get())
            tmp = SETTINGS_FILE.with_suffix(".json.tmp")
            tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            os.replace(tmp, SETTINGS_FILE)
        except Exception as e:
            self.append(f"保存搜索设置失败：{type(e).__name__}: {e}")

    # --- 搜索结果的多选（复选框）与批量操作 ---------------------------------
    def _on_search_tree_click(self, event):
        """点第 1 列的方框切换勾选；点其它列保持默认的整行高亮选择。"""
        if self.search_tree.identify_region(event.x, event.y) not in ("cell", "tree"):
            return None
        if self.search_tree.identify_column(event.x) != "#1":
            return None
        iid = self.search_tree.identify_row(event.y)
        if not iid:
            return None
        self.toggle_search_check(iid)
        return "break"

    def set_search_check(self, iid, checked=True):
        """设置某一行勾选状态（行号存在 _search_row_index 里，不靠解析展示文字）。"""
        if not iid or not self.search_tree.exists(iid):
            return
        if checked:
            self._search_checked.add(iid)
        else:
            self._search_checked.discard(iid)
        try:
            self.search_tree.set(iid, "pick", "☑" if checked else "☐")
            # 勾选的行整行绿色加粗，勾没勾一眼可辨。「offcampus」排在后面，
            # 所以非本校校区的红字仍然压得住颜色（红字 + 加粗）。
            tags = [t for t in (self.search_tree.item(iid, "tags") or ()) if t != "checked"]
            if checked:
                tags.insert(0, "checked")
            # ttk.Treeview.item() 只有 (item, option)：写值必须走关键字参数。
            self.search_tree.item(iid, tags=tags)
        except tk.TclError:
            pass
        self.update_pick_hint()

    def toggle_search_check(self, iid=None):
        """切换勾选：给了 iid 就切换那一行，否则切换当前高亮选中的行（空格键）。"""
        targets = [iid] if iid else list(self.search_tree.selection())
        for item in targets:
            self.set_search_check(item, item not in self._search_checked)

    def check_all_search(self, checked=True):
        for item in self.search_tree.get_children():
            self.set_search_check(item, bool(checked))

    def toggle_select_all(self):
        """全选/全不选：表头点「选」或按 Ctrl+A 都走这里（已全勾就一键清空）。"""
        items = self.search_tree.get_children()
        if not items:
            return
        every = all(item in self._search_checked for item in items)
        self.check_all_search(not every)

    def _on_search_select_all_key(self, _event=None):
        self.toggle_select_all()
        return "break"

    def invert_search_checks(self):
        for item in self.search_tree.get_children():
            self.set_search_check(item, item not in self._search_checked)

    def checked_search_indices(self):
        """要操作的行序号：优先用方框勾选的，一行都没勾就用高亮选中的行（两种操作都支持）。"""
        picked = sorted({self._search_row_index[iid] for iid in self.search_tree.get_children()
                         if iid in self._search_checked and iid in self._search_row_index})
        if picked:
            return picked
        fallback = []
        for iid in self.search_tree.selection():
            if iid in self._search_row_index:
                fallback.append(self._search_row_index[iid])
                continue
            try:      # 兜底：外部塞进表里的行，按「序号」列解析
                fallback.append(int(self.search_tree.item(iid, "values")[1]) - 1)
            except (IndexError, TypeError, ValueError):
                continue
        return sorted(set(fallback))

    def update_pick_hint(self):
        if not hasattr(self, "pick_hint"):
            return
        count = len(self._search_checked)
        self.pick_hint.configure(
            text=(f"已勾选 {count} 个教学班 → 可批量加入抢课 / 批量抢课" if count
                  else "点第一列方框可勾选多行 → 批量加入抢课 / 批量抢课"))

    def fill_search_tree(self, entries) -> int:
        """重建搜索结果表，返回其中非本校校区（教工路）的条数。

        行号记在 _search_row_index（iid → 序号），不解析展示文字：加了「选」「校区」
        两列以后列顺序变了也不会取错行。
        """
        for item in self.search_tree.get_children():
            self.search_tree.delete(item)
        self._search_row_index.clear()
        self._search_checked.clear()
        off_campus = 0
        for index, entry in enumerate(entries):
            campus = normalize_campus_name(entry.get("campus")) or CAMPUS_LOCAL
            platform = str(entry.get("platform") or "").strip()
            off = is_off_campus(campus)
            if off:
                off_campus += 1
            if off:
                campus_text = f"⚠{campus}"
            elif platform:
                campus_text = "网课"
            else:
                campus_text = campus
            time_text = entry.get("time", "") or "—"
            if platform:
                time_text = f"【{platform}】{time_text}"
            iid = self.search_tree.insert("", "end", values=(
                "☐", index + 1, entry.get("name", ""), entry.get("credit") or "—",
                entry.get("class", ""), campus_text, time_text,
                entry.get("teacher", "") or "—", entry.get("status", "")),
                tags=("offcampus",) if off else ())
            self._search_row_index[iid] = index
        self.update_pick_hint()
        return off_campus

    def _search_entry(self, index):
        """取搜索结果第 index 行：(本地行字典, None) 或 (None, Jxb)；越界返回 None。"""
        result = getattr(self, "search_result", None)
        if not result or index is None or index < 0:
            return None
        if isinstance(result, LocalSearchResult):
            rows = result.rows or []
            return (rows[index], None) if index < len(rows) else None
        jxbs = list(getattr(result, "jxbs", None) or [])
        return (None, jxbs[index]) if index < len(jxbs) else None

    def _preset_from_search_row(self, row, jxb, mode, result):
        """把一行搜索结果转成 (关键字, 类别, 详情)。

        关键字就是抢课时真正发给教务的查询词：默认锁定教学班名（含校区前缀），
        模式为「课程名」/「课程代码」时才退化成按名/按号抢。
        """
        if row is not None:
            category = course_category(row)
            detail = self._catalog_row_to_detail(row.get("teaching_class") or row.get("jxb_id"), row)
            if mode == "课程名":
                keyword = row.get("course_name") or row.get("keyword") or ""
            elif mode == "课程代码":
                keyword = row.get("kch") or row.get("course_code") or ""
            else:
                keyword = (row.get("teaching_class") or row.get("jxb_id")
                           or row.get("course_name") or "")
            return str(keyword or "").strip(), category, detail
        course = getattr(result, "course", None)
        tab_name = getattr(course, "tab_name", "") or ""
        category = ("体育" if "体育" in tab_name else "英语" if "英语" in tab_name
                    else "通识" if "通识" in tab_name else "OTHER")
        if mode == "课程名":
            keyword = getattr(course, "kcmc", "") or ""
        elif mode == "课程代码":
            keyword = getattr(course, "kch_id", "") or ""
        else:
            keyword = getattr(jxb, "jxbmc", "") or getattr(jxb, "jxb_id", "") or ""
        return str(keyword or "").strip(), category, self.detail_from_fetch(result, jxb, category)

    def build_all(self):
        f = self.all_tab
        ttk.Label(f, text="全部课程", font=("Microsoft YaHei UI", 12, "bold")).pack(anchor="w")
        ttk.Label(f, text="直连教务接口分页拉取全部教学班（约 1900+ 条），可导出 Excel。", wraplength=900).pack(anchor="w", pady=8)
        btnrow = ttk.Frame(f)
        btnrow.pack(anchor="w", pady=8)
        ttk.Button(btnrow, text="拉取全部课程", style="Accent.TButton", command=self.pull_all_courses_ui).pack(side="left")
        ttk.Button(btnrow, text="导出 Excel", command=self.export_all_courses_excel).pack(side="left", padx=10)
        ttk.Label(btnrow, text="⚠ 教工路 = 非你所在校区（下沙），需跨校区通勤", style="Muted.TLabel").pack(side="left", padx=6)
        cols = ("idx", "name", "kch", "jxb", "campus", "time", "room", "sel")
        self.all_tree = ttk.Treeview(f, columns=cols, show="headings", height=12)
        for c, t, w in [("idx", "#", 44), ("name", "课程名", 190), ("kch", "课程号", 95),
                        ("jxb", "教学班", 165), ("campus", "校区", 74), ("time", "上课时间", 240),
                        ("room", "教室", 120), ("sel", "已选/容量", 86)]:
            self.all_tree.heading(c, text=t)
            self.all_tree.column(c, width=w, anchor="center" if c in ("idx", "campus", "sel") else "w")
        self.all_tree.tag_configure("offcampus", foreground="#b91c1c")
        self.all_tree.pack(fill="both", expand=True, pady=8)

    def build_selected(self):
        f = self.selected_tab
        f.columnconfigure(0, weight=1)
        f.rowconfigure(3, weight=3, minsize=int(300 * DPI_SCALE))   # 课表行：可随窗口放大
        f.rowconfigure(4, weight=2, minsize=int(150 * DPI_SCALE))   # 明细表行

        top = ttk.Frame(f)
        top.grid(row=0, column=0, sticky="ew")
        self.completion_var = tk.StringVar(value="选课未完成")
        ttk.Label(top, textvariable=self.completion_var,
                  font=("Microsoft YaHei UI", 16, "bold")).pack(side="left")
        self.selected_summary_var = tk.StringVar(value="等待刷新教务系统")
        ttk.Label(top, textvariable=self.selected_summary_var, style="Muted.TLabel").pack(side="left", padx=16)
        ttk.Button(top, text="刷新已选课程", style="Accent.TButton",
                   command=self.refresh_selected_courses).pack(side="right")
        ttk.Button(top, text="全屏 / 还原", command=self.toggle_fullscreen).pack(side="right", padx=8)
        refresh_row = ttk.Frame(f)
        refresh_row.grid(row=1, column=0, sticky="ew")
        self.selected_refresh_var = tk.StringVar(value="已选数据：尚未刷新")
        ttk.Label(refresh_row, textvariable=self.selected_refresh_var, style="Muted.TLabel").pack(side="left")
        # 实时对齐：空闲时自动刷新教务已选（抢课运行期间不并发发请求）
        self.auto_refresh_var = tk.BooleanVar(value=True)
        self.auto_refresh_btn = ttk.Button(refresh_row, command=lambda: self.toggle_option(
            self.auto_refresh_var, self.auto_refresh_btn, "自动对齐已选（每60秒）"))
        self.auto_refresh_btn.pack(side="right")
        self.refresh_toggle(self.auto_refresh_var, self.auto_refresh_btn, "自动对齐已选（每60秒）")

        progress = ttk.Frame(f)
        progress.grid(row=2, column=0, sticky="ew", pady=(6, 6))
        self.progress_widgets = {}
        for category, label, target in (("EN", "英语", 1), ("PE", "体育", 1), ("GE", "待抢通识", 2)):
            card = ttk.Frame(progress, style="Card.TFrame", padding=8)
            card.pack(side="left", fill="x", expand=True, padx=(0, 8))
            var = tk.StringVar(value=f"{label} 0/{target}")
            ttk.Label(card, textvariable=var, style="Card.TLabel").pack(anchor="w")
            bar = ttk.Progressbar(card, maximum=target, value=0)
            bar.pack(fill="x", pady=(5, 0))
            self.progress_widgets[category] = (var, bar, label, target)

        tt_host = ttk.Frame(f)
        tt_host.grid(row=3, column=0, sticky="nsew", pady=(2, 4))
        tt_host.columnconfigure(0, weight=1)
        tt_host.rowconfigure(0, weight=1)
        self.timetable_canvas = tk.Canvas(tt_host, height=int(300 * DPI_SCALE), bg="#ffffff",
                                          highlightthickness=1, highlightbackground=SOFT)
        self.timetable_canvas.grid(row=0, column=0, sticky="nsew")
        self.timetable_canvas.bind("<Configure>", lambda _e: self.draw_timetable())
        ttk.Label(tt_host, text="课表说明：同一时间段出现多个色块时，课程会自动并排显示；这代表真实时间重叠，不会再互相覆盖。窗口最大化或全屏后，课表会随空间放大。", style="Muted.TLabel").grid(row=1, column=0, sticky="w", pady=(0, 2))

        cols = ("name", "code", "class", "time", "teacher", "room", "status")
        table_frame = ttk.Frame(f)
        table_frame.grid(row=4, column=0, sticky="nsew", pady=(0, 4))
        table_frame.columnconfigure(0, weight=1)
        table_frame.rowconfigure(0, weight=1)
        self.selected_tree = ttk.Treeview(table_frame, columns=cols, show="headings", height=6)
        for col, title, width in [
            ("name", "课程名", 150), ("code", "课程号", 80), ("class", "教学班", 190),
            ("time", "上课时间", 230), ("teacher", "教师", 130), ("room", "教室", 100),
            ("status", "状态", 80),
        ]:
            self.selected_tree.heading(col, text=title)
            self.selected_tree.column(col, width=width, anchor="w")
        self.selected_tree.grid(row=0, column=0, sticky="nsew")
        selected_scroll = ttk.Scrollbar(table_frame, orient="vertical", command=self.selected_tree.yview)
        selected_scroll.grid(row=0, column=1, sticky="ns")
        self.selected_tree.configure(yscrollcommand=selected_scroll.set)
        self.update_selected_dashboard()

    def load_system_selected_cache(self):
        self.system_selected_count = 0
        self.system_selected_credits = ""
        try:
            data = json.loads(SYSTEM_SELECTED_FILE.read_text(encoding="utf-8"))
            self.system_selected_count = int(data.get("count", 0))
            self.system_selected_credits = str(data.get("credits", ""))
            return data.get("courses", []) if isinstance(data.get("courses"), list) else []
        except Exception:
            return []

    def load_course_catalog(self):
        try:
            data = json.loads(COURSE_CATALOG_FILE.read_text(encoding="utf-8"))
            if not isinstance(data, dict):
                return {}
            # 旧缓存可能只有基础字段；用随程序保存的详细课程快照补齐分类/余量。
            detailed_file = DATA_DIR / "all_courses_detailed.json"
            if detailed_file.exists() and any("kclbmc" not in row for row in data.values()):
                detailed = json.loads(detailed_file.read_text(encoding="utf-8"))
                by_class = {row.get("jxbmc"): row for row in detailed if row.get("jxbmc")}
                changed = False
                for teaching_class, row in data.items():
                    extra = by_class.get(teaching_class)
                    if not extra:
                        continue
                    for target, source in (("kclbmc", "kclbmc"), ("kcxzmc", "kcxzmc"),
                                           ("yixuan", "yxzrs"), ("rongliang", "jxbrl"),
                                           ("course_group", "kclbmc"), ("teacher", "jsxx"),
                                           ("xf", "xf")):
                        if not row.get(target) and extra.get(source) not in (None, ""):
                            row[target] = extra.get(source)
                            changed = True
                if changed:
                    tmp = COURSE_CATALOG_FILE.with_suffix(".json.tmp")
                    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
                    os.replace(tmp, COURSE_CATALOG_FILE)
            return data
        except Exception:
            return {}

    def save_course_catalog(self, rows):
        catalog = {}
        for row in rows:
            teaching_class = row.get("teaching_class")
            if teaching_class:
                catalog[teaching_class] = {
                    "course_name": row.get("course_name") or row.get("keyword", ""),
                    "course_code": row.get("kch", ""), "time": row.get("time", ""),
                    "room": row.get("room", ""), "teacher": row.get("teacher", ""),
                    "course_tab": row.get("course_tab", ""),
                    "course_belong": row.get("course_belong", ""),
                    "course_group": row.get("course_group", "") or row.get("kclbmc", ""),
                    "kclbmc": row.get("kclbmc", ""),
                    "kcxzmc": row.get("kcxzmc", ""),
                    "xf": row.get("xf", ""),            # 学分（旧版漏存，导致缓存里没有学分）
                    "yixuan": row.get("yixuan", ""), "rongliang": row.get("rongliang", ""),
                    "xqumc": row.get("xqumc", ""),
                    "do_id": row.get("do_id", ""), "jxb_id": row.get("jxb_id", ""),
                }
        tmp = COURSE_CATALOG_FILE.with_suffix(".json.tmp")
        tmp.write_text(json.dumps(catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        os.replace(tmp, COURSE_CATALOG_FILE)
        self.course_catalog = catalog
        self._catalog_core_index = None  # 课程库变了 → 核心教学班号索引重建

    def save_system_selected_cache(self, result):
        tmp = SYSTEM_SELECTED_FILE.with_suffix(".json.tmp")
        tmp.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        os.replace(tmp, SYSTEM_SELECTED_FILE)

    def catalog_entry(self, teaching_class: str) -> dict:
        """按教学班名取课程库详情：**带前缀和不带前缀的写法都认**。

        教务同一门课有时给「国情课(2026-2027-1)-GENEML053-01」、有时给
        「(2026-2027-1)-GENEML053-01」，而课程库里只存了其中一种写法。
        只按整串取会让另一条取不到详情 → 教学班名/时间/学分空着，GE/EN/PE 归类还会掉
        （进度显示 1/2 而不是 2/2，抢课目标跟着算错）。2026-09-17 修这个去重时实测踩到。
        """
        text = str(teaching_class or "").strip()
        entry = self.course_catalog.get(text)
        if entry:
            return entry
        index = getattr(self, "_catalog_core_index", None)
        if index is None:
            index = {}
            for key, value in (self.course_catalog or {}).items():
                core = teaching_class_core(key)
                if core and core not in index:
                    index[core] = value
            self._catalog_core_index = index
        return index.get(teaching_class_core(text)) or {}

    def combined_selected_courses(self):
        """已选课程 = 教务已选 + 本程序抢到（同一门课只出现一次）。

        去重按**核心教学班号**：教务那边同一门课有时带前缀、有时不带
        （国情课(2026-2027-1)-GENEML053-01 vs (2026-2027-1)-GENEML053-01），
        比整串字符串会把同一个班当成两门课——2026-09-17 用户就遇到
        "已选列表里扒出两个中国城市经济与发展"。
        """
        courses = []
        seen = set()

        def add(course):
            core = teaching_class_core(course.get("teaching_class"))
            if core:
                key = core.lower()
            else:
                key = "{}|{}".format(str(course.get("course_code") or "").strip().upper(),
                                     str(course.get("course_name") or "").strip())
            if key.strip("|") and key in seen:
                return False
            seen.add(key)
            courses.append(course)
            return True

        for item in self.system_selected:
            course = dict(item)
            detail = self.catalog_entry(course.get("teaching_class"))
            for field in ("course_name", "course_code", "time", "room", "teacher"):
                if not course.get(field) and detail.get(field):
                    course[field] = detail[field]
            add(course)
        for item in self.state_store.selected():
            merged = dict(item)
            keyword = item.get("keyword", "")
            detail = (self.preset_details.get(keyword) or [{}])[0]
            merged.setdefault("course_name", detail.get("course_name", ""))
            merged.setdefault("course_code", detail.get("code", ""))
            merged.setdefault("time", detail.get("time", ""))
            merged.setdefault("room", detail.get("room", ""))
            merged["status"] = "本次已抢到"
            if not add(merged):
                continue          # 教务已选里已有同一个班（前缀不同也算）→ 不重复列一条
        return courses

    def course_category_code(self, course) -> str:
        """把一条课程映射成 GE/EN/PE/OTHER。

        判据优先级：本地「全部课程」缓存里的选课 Tab（教务口径）→ 记录自带的 category。
        这样"教务系统里已经选上的通识/英语/体育"才会被算进名额进度。
        """
        own = str(course.get("category") or "").upper()
        detail = self.catalog_entry(course.get("teaching_class") or "")
        row = {
            "kklxdm": detail.get("kklxdm", ""),
            "course_tab": detail.get("course_tab") or detail.get("category") or "",
            "kclbmc": detail.get("kclbmc", ""),
            "course_belong": detail.get("course_belong", ""),
            "course_name": course.get("course_name") or detail.get("course_name") or "",
        }
        mapped = {"英语": "EN", "体育": "PE", "通识": "GE"}.get(course_category(row))
        if mapped:
            return mapped
        return own if own in {"GE", "EN", "PE"} else "OTHER"

    def annotated_selected_courses(self):
        """已选课程（教务 + 本程序）并集，并补上 GE/EN/PE 归类，供进度统计使用。"""
        rows = []
        for course in self.combined_selected_courses():
            item = dict(course)
            item["category"] = self.course_category_code(course)
            item.setdefault("keyword", course.get("course_name") or course.get("teaching_class") or "")
            rows.append(item)
        return rows

    def progress_snapshot(self):
        """进度快照：以"教务已选 + 本程序抢到"的并集为准（实时对齐，不是只看本地记录）。"""
        snapshot = completion_snapshot(self.annotated_selected_courses(), self.state_store.targets)
        snapshot["pending"] = self.state_store.pending()
        return snapshot

    def occupied_times(self):
        """已选课程占用的上课时间文本（供时间段分配用）。"""
        return [course.get("time") for course in self.annotated_selected_courses() if course.get("time")]

    def update_selected_dashboard(self, snapshot=None):
        snapshot = snapshot or self.progress_snapshot()
        self.completion_var.set("✓ 选课已完成" if snapshot.get("completed") else "○ 选课尚未完成")
        for category, (var, bar, label, target) in self.progress_widgets.items():
            done, _goal = snapshot.get("progress", {}).get(category, (0, target))
            var.set(f"{label} {done}/{target}")
            bar.configure(value=done)
        credits = f"，{self.system_selected_credits} 学分" if self.system_selected_credits else ""
        self.selected_summary_var.set(f"教务系统已选 {self.system_selected_count} 门{credits}；待选 {len(snapshot.get('pending', []))} 项")
        if hasattr(self, "selected_refresh_var"):
            when = self.last_selected_refresh.strftime("%H:%M:%S") if self.last_selected_refresh else "尚未刷新"
            mode = "自动对齐中" if getattr(self, "auto_refresh_var", None) and self.auto_refresh_var.get() else "自动对齐已关闭"
            if mode == "自动对齐中":
                # 如实写出"为什么还没刷新"，不然用户会以为功能没了（2026-09-17 就是这么误会过）
                block = self.auto_refresh_block_reason()
                if block:
                    mode += f"·{block}"
            self.selected_refresh_var.set(f"已选数据：{when}（{mode}，每 {self.auto_refresh_interval} 秒）")
        for row in self.selected_tree.get_children():
            self.selected_tree.delete(row)
        for course in self.combined_selected_courses():
            self.selected_tree.insert("", "end", values=(
                course.get("course_name", ""), course.get("course_code", ""),
                course.get("teaching_class", ""), course.get("time", ""),
                course.get("teacher", ""), course.get("room", ""),
                course.get("status", "已选上"),
            ))
        self.draw_timetable()

    def draw_timetable(self):
        if not hasattr(self, "timetable_canvas"):
            return
        canvas = self.timetable_canvas
        canvas.delete("all")
        width = max(canvas.winfo_width(), 640)
        height = max(canvas.winfo_height(), int(200 * DPI_SCALE))
        left, top = 42, 24
        col_w = (width - left - 6) / 5
        row_h = (height - top - 4) / 12
        days = ["周一", "周二", "周三", "周四", "周五"]
        for i, day in enumerate(days):
            x = left + i * col_w
            canvas.create_text(x + col_w / 2, 12, text=day, fill=TEXT, font=("Microsoft YaHei UI", 9, "bold"))
        for period in range(1, 13):
            y = top + (period - 1) * row_h
            canvas.create_text(20, y + row_h / 2, text=str(period), fill=MUTED, font=("Microsoft YaHei UI", 8))
            canvas.create_line(left, y, width - 4, y, fill=SOFT)
        for i in range(6):
            x = left + i * col_w
            canvas.create_line(x, top, x, height - 4, fill=SOFT)
        laid_out = layout_timetable_slots(timetable_slots(self.combined_selected_courses()))
        colors = ["#dff7ef", "#e8efff", "#fff1d6", "#f4e8ff", "#ffe7ec"]
        for slot in laid_out:
            lane_count = max(1, slot.get("lanes", 1))
            lane_w = (col_w - 6) / lane_count
            x1 = left + (slot["day"] - 1) * col_w + 2 + slot.get("lane", 0) * lane_w
            x2 = x1 + lane_w - 3
            y1 = top + (slot["start"] - 1) * row_h + 1
            y2 = top + slot["end"] * row_h - 1
            color = colors[sum(ord(ch) for ch in slot["course_name"]) % len(colors)]
            canvas.create_rectangle(x1, y1, x2, y2, fill=color, outline="#9fb8b0")
            font_size = max(7, min(12, int(lane_w / 30), int(row_h / 1.9)))
            canvas.create_text((x1 + x2) / 2, (y1 + y2) / 2,
                               text=f"{slot['course_name']}\n{slot['weeks']}", width=max(40, lane_w - 8),
                               fill=TEXT, font=("Microsoft YaHei UI", font_size), justify="center")

    def toggle_fullscreen(self):
        current = bool(self.attributes("-fullscreen"))
        self.attributes("-fullscreen", not current)

    def refresh_selected_courses(self):
        base_url, cookie = self.base_var.get(), self.cookie_text()
        def work():
            client = make_client(base_url, cookie, self.thread_log)
            init_until_open(client, self.thread_log, once=True)
            return ("selected_courses", client.fetch_selected_courses(enrich=True))
        self.run_thread(work, "刷新已选课程中……")

    def _cancel_timers(self):
        """退出前取消未执行的定时回调，避免 Tk 报 'invalid command name ...poll'。"""
        for attr in ("_poll_after_id", "_tick_after_id"):
            after_id = getattr(self, attr, None)
            if after_id:
                try:
                    self.after_cancel(after_id)
                except Exception:
                    pass
                setattr(self, attr, None)

    def destroy(self):
        self._cancel_timers()
        super().destroy()

    def auto_refresh_tick(self):
        """定时器：空闲时自动把"已选课程/进度"对齐到教务真相。"""
        try:
            self.maybe_auto_refresh()
        except Exception as e:
            self.append(f"自动对齐检查异常：{type(e).__name__}: {e}")
        try:
            self._tick_after_id = self.after(5000, self.auto_refresh_tick)
        except tk.TclError:
            self._tick_after_id = None

    def auto_refresh_block_reason(self) -> str:
        """「自动对齐已选」当前跑不了的原因（能跑就返回空串）。

        必须如实说出来：2026-09-17 用户重启软件后，Cookie 框是空的 → 自动对齐一直静默跳过、
        日志里也再没有「已选课程刷新完成」，用户以为这个功能被删了。定时器其实一直在跑，
        只是没 Cookie 时什么都不做、也不吭声。
        """
        if self.auto_refresh_var is None or not self.auto_refresh_var.get():
            return "开关已关"
        if not self.cookie_text():
            return "缺 Cookie，点一下「抓取 Cookie」就开始"
        if self.worker is not None and self.worker.is_alive():
            return "有任务在跑，等它结束"
        return ""

    def _note_auto_refresh_block(self, reason: str):
        """把"自动对齐为什么没执行"记一次日志（同一条原因只记一次，不刷屏）。"""
        if reason == getattr(self, "_auto_refresh_block_logged", None):
            return
        self._auto_refresh_block_logged = reason
        if reason:
            self.append(f"自动对齐已选：暂时跳过（{reason}）")

    def maybe_auto_refresh(self, force: bool = False) -> bool:
        """需要时刷新已选课程。

        抢课/其它任务运行时不并发发请求（同一个 worker），等它跑完再对齐——
        抢课过程中进度按"本程序抢到"的本地结果即时更新，跑完后立刻用教务真相兜一次。
        """
        reason = self.auto_refresh_block_reason()
        if reason:
            if force or self.last_selected_refresh is None:
                self._note_auto_refresh_block(reason)
            return False
        if (not force and self.last_selected_refresh is not None
                and (datetime.now() - self.last_selected_refresh).total_seconds() < self.auto_refresh_interval):
            return False
        self._note_auto_refresh_block("")      # 能跑了 → 清掉"上次没跑"的记录
        self.refresh_selected_courses()
        return True

    def request_selected_refresh_soon(self, delay_ms: int = 300):
        """等当前 worker 退出后再刷新（避免和正在结束的任务抢同一个 worker）。"""
        self.after(delay_ms, lambda: self.maybe_auto_refresh(force=True))

    def reload_pending_presets(self):
        self.preset_list.delete(0, "end")
        self.load_presets_from_file()
        self.refresh_preset_display()

    def on_course_success(self, raw_preset, info):
        # 待选列表控件里存的是「展示行」（类别｜课程名｜教学班｜…），
        # 直接落盘会写出上百字的怪关键字，而且类别会变成 OTHER 导致
        # EN/GE/PE 名额上限统计失灵。统一先规范化成 GE:xxx / EN:xxx 形式。
        raw_preset = canonical_preset(raw_preset)
        keyword, _category = parse_preset(raw_preset)
        detail = (self.preset_details.get(keyword) or [{}])[0]
        info = dict(info)
        info.setdefault("course_name", detail.get("course_name", ""))
        info.setdefault("course_code", detail.get("code", ""))
        info.setdefault("room", detail.get("room", ""))
        snapshot = self.state_store.mark_completed(raw_preset, info, source="assistant")
        self.q.put(("course_success", (raw_preset, snapshot)))
        return snapshot

    def refresh_toggle(self, var, btn, label):
        on = bool(var.get())
        btn.configure(text=("✓ 已开启：" if on else "○ 已关闭：") + label, style="On.TButton" if on else "Off.TButton")

    def toggle_option(self, var, btn, label):
        var.set(not bool(var.get()))
        self.refresh_toggle(var, btn, label)
        if hasattr(self, "search_filters"):
            self.save_runtime_settings()

    def append(self, msg):
        text = str(msg)
        try:
            with self.log_file.open("a", encoding="utf-8") as fp:
                fp.write(f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] {text}\n")
            self._enforce_log_budget()
        except Exception:
            pass
        self.log.configure(state="normal")
        self.log.insert("end", text + "\n")
        self.log.see("end")
        self.log.configure(state="disabled")

    def _migrate_data_file(self, name):
        """首次升级时把旧项目目录中的用户数据迁移到 AppData。"""
        target = DATA_DIR / name
        legacy = PROJECT_DIR / name
        if not target.exists() and legacy.exists():
            try:
                target.write_bytes(legacy.read_bytes())
            except Exception:
                pass

    def _migrate_legacy_logs(self):
        """把旧项目 logs 中的历史日志复制到统一日志目录（不删除原文件）。"""
        legacy_dir = PROJECT_DIR / "logs"
        if not legacy_dir.exists() or any(LOG_DIR.glob("*.log")):
            return
        try:
            for path in legacy_dir.glob("*.log"):
                target = LOG_DIR / f"legacy_{path.name}"
                target.write_bytes(path.read_bytes())
        except Exception:
            pass

    def _enforce_log_budget(self):
        """按文件轮转并清理旧日志，日志目录总大小硬限制为 1 MiB。"""
        try:
            if self.log_file.exists() and self.log_file.stat().st_size > LOG_ROTATE_SIZE:
                stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
                self.log_file.replace(LOG_DIR / f"app_{stamp}.log")
                self.log_file.touch()
            files = sorted(LOG_DIR.glob("*.log"), key=lambda p: p.stat().st_mtime, reverse=True)
            total = 0
            for path in files:
                size = path.stat().st_size
                if total + size <= LOG_BUDGET:
                    total += size
                    continue
                path.unlink(missing_ok=True)
        except Exception:
            pass

    def open_log_folder(self):
        try:
            os.startfile(str(LOG_DIR))
        except Exception as e:
            messagebox.showerror("打开失败", f"{type(e).__name__}: {e}")

    def thread_log(self, msg):
        self.q.put(("log", str(msg)))

    def cookie_text(self):
        return self.cookie_box.get("1.0", "end").strip()

    def run_thread(self, fn, status="运行中……"):
        if self.worker and self.worker.is_alive():
            messagebox.showinfo("正在运行", "当前任务还没结束。")
            return
        self.status.set(status)
        self.update_task_progress(f"▶ 开始：{status}", "运行中")
        def task():
            try:
                self.q.put(fn())
            except Exception as e:
                self.q.put(("error", f"{type(e).__name__}: {e}\n{traceback.format_exc(limit=4)}"))
        self.worker = threading.Thread(target=task, daemon=True)
        self.worker.start()

    def on_close(self):
        self.save_runtime_settings()
        self.destroy()

    def load_runtime_settings(self):
        try:
            data = json.loads(SETTINGS_FILE.read_text(encoding="utf-8")) if SETTINGS_FILE.exists() else {}
            for name, var in (("base_url", self.base_var), ("interval", self.interval_var),
                              ("max_success", self.max_success_var), ("max_rounds", self.max_rounds_var),
                              ("selector", self.selector_var), ("start_time", self.start_time_var),
                              ("focus_count", self.focus_count_var), ("full_wait", self.full_wait_var)):
                if name in data:
                    var.set(data[name])
            for name, var in (("retry", self.retry_var), ("watch_full", self.watch_full_var),
                              ("auto_after_init", self.auto_after_init_var), ("search_loop", self.search_loop_var),
                              ("auto_refresh_selected", self.auto_refresh_var),
                              ("avoid_conflict", self.avoid_conflict_var)):
                if name in data:
                    var.set(bool(data[name]))
            self.refresh_toggle(self.retry_var, self.retry_btn, "频率过高自动重试")
            self.refresh_toggle(self.watch_full_var, self.watch_full_btn, "满课蹲守退课")
            self.refresh_toggle(self.auto_after_init_var, self.auto_after_init_btn, "init成功后自动抢课")
            self.refresh_toggle(self.search_loop_var, self.search_loop_btn, "循环抢课（每1.5秒重试）")
            self.refresh_toggle(self.auto_refresh_var, self.auto_refresh_btn, "自动对齐已选（每60秒）")
            self.refresh_toggle(self.avoid_conflict_var, self.avoid_conflict_btn, "自动避开时间冲突")
        except Exception as e:
            self.append(f"读取软件设置失败：{type(e).__name__}: {e}")

    def save_runtime_settings(self):
        try:
            data = {}
            if SETTINGS_FILE.exists():
                data = json.loads(SETTINGS_FILE.read_text(encoding="utf-8"))
            data.update({
                "base_url": self.base_var.get(), "interval": self.interval_var.get(),
                "max_success": self.max_success_var.get(), "max_rounds": self.max_rounds_var.get(),
                "selector": self.selector_var.get(), "start_time": self.start_time_var.get(),
                "retry": bool(self.retry_var.get()), "watch_full": bool(self.watch_full_var.get()),
                "auto_after_init": bool(self.auto_after_init_var.get()), "search_loop": bool(self.search_loop_var.get()),
                "auto_refresh_selected": bool(self.auto_refresh_var.get()) if self.auto_refresh_var else True,
                "avoid_conflict": bool(self.avoid_conflict_var.get()) if self.avoid_conflict_var else True,
                "focus_count": self.focus_count_var.get(), "full_wait": self.full_wait_var.get(),
            })
            data["search_filters"] = self.search_filters
            data["search_keyword"] = self.search_var.get()
            data["search_mode"] = self.search_mode.get()
            tmp = SETTINGS_FILE.with_suffix(".json.tmp")
            tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            os.replace(tmp, SETTINGS_FILE)
        except Exception as e:
            self.append(f"保存软件设置失败：{type(e).__name__}: {e}")

    def update_task_progress(self, message, state=None):
        """更新日志下方的任务进度框，并在日志中保留详细步骤。"""
        if not hasattr(self, "task_progress"):
            return
        if state:
            self.task_progress_label.configure(text=f"当前状态：{state}")
        if state == "运行中":
            self._progress_finished = False
            self.task_progress.start(12)

    def finish_task_progress(self, label="就绪"):
        self.task_progress.stop()
        self._progress_finished = True
        self.update_task_progress(f"■ 任务{label}", label)

    def settle_task_progress(self):
        """兜底收尾：任务线程已结束、队列也已排空，但界面还停在“运行中”时就地收尾。

        抓 Cookie / 刷新人数 / 拉取全部课程 / 刷新已选课程 这些分支以前没有调用
        finish_task_progress，任务跑完进度框会一直转、状态一直是“运行中”。
        """
        if getattr(self, "_progress_finished", True):
            return
        if self.worker is not None and self.worker.is_alive():
            return
        if not self.q.empty():
            return
        self.finish_task_progress("已完成")

    def open_browser(self):
        try:
            launch_login_browser(self.base_var.get(), self.append)
        except Exception as e:
            messagebox.showerror("错误", str(e))

    def close_browser(self):
        try:
            close_login_browser(self.append)
        except Exception as e:
            messagebox.showerror("关闭失败", str(e))

    def capture_cookie(self):
        self.run_thread(lambda: ("cookie", capture_cookie_from_edge(self.thread_log)), "抓取 Cookie 中……")

    def verify_cookie(self):
        base_url, cookie = self.base_var.get(), self.cookie_text()
        def work():
            status = verify_cookie_full(base_url, cookie, self.thread_log)
            return ("status", status)
        self.run_thread(work, "验证 Cookie 中……")

    def copy_cookie(self):
        cookie = self.cookie_text()
        if not cookie:
            messagebox.showwarning("没有 Cookie", "请先自动抓取 Cookie，或手动粘贴 Cookie。")
            return
        self.clipboard_clear()
        self.clipboard_append(cookie)
        self.update()
        self.append("完整 Cookie 已复制到剪贴板，可粘贴到外部样板软件。")
        self.status.set("Cookie 已复制到剪贴板")

    def check_cookie_login(self):
        self.verify_cookie()

    def add_preset(self):
        v = self.preset_var.get().strip()
        if v:
            self.preset_list.insert("end", v)
            self.preset_var.set("")
            self.refresh_preset_display()
            self.persist_presets()

    def move_preset(self, delta):
        sel = self.preset_list.curselection()
        if not sel:
            return
        i = sel[0]
        j = i + delta
        if j < 0 or j >= self.preset_list.size():
            return
        v = self.preset_list.get(i)
        self.preset_list.delete(i)
        self.preset_list.insert(j, v)
        self.preset_list.selection_set(j)
        self.refresh_preset_display()
        self.persist_presets()

    def delete_preset(self):
        for i in reversed(self.preset_list.curselection()):
            self.preset_list.delete(i)
        self.refresh_preset_display()
        self.persist_presets()

    def clear_presets(self):
        if self.preset_list.size() == 0:
            return
        if messagebox.askyesno("确认", "确认清空全部预设课程吗？"):
            count = self.preset_list.size()
            self.preset_list.delete(0, "end")
            self.refresh_preset_display()
            self.persist_presets()
            # 这个操作会改写待选文件，必须留痕（否则事后查日志看不出列表为什么空了）。
            self.append(f"已清空待选列表（{count} 条）。")
            self.status.set("待选列表已清空")

    def presets(self):
        return [self.preset_list.get(i) for i in range(self.preset_list.size())]

    def persist_presets(self):
        self.state_store.replace_pending([canonical_preset(raw) for raw in self.presets()])

    def load_preset_details(self):
        try:
            self.preset_details = json.loads(PRESET_DETAILS_FILE.read_text(encoding="utf-8")) if PRESET_DETAILS_FILE.exists() else {}
        except Exception:
            self.preset_details = {}
        self.preset_details_dirty = False

    def save_preset_details(self):
        """把预设详情写回 course_preset_details.json（原子替换，失败不影响界面）。"""
        try:
            data = {key: [{k: v for k, v in (item or {}).items() if not str(k).startswith("_")}
                          for item in (value if isinstance(value, list) else [value])]
                    for key, value in self.preset_details.items()}
            tmp = PRESET_DETAILS_FILE.with_suffix(".json.tmp")
            tmp.write_text(json.dumps(data, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
            os.replace(tmp, PRESET_DETAILS_FILE)
            self.preset_details_dirty = False
            return True
        except Exception as e:
            self.append(f"保存预设详情失败：{type(e).__name__}: {e}")
            return False

    def merge_preset_detail(self, keyword, detail) -> bool:
        """把新拿到的详情并进已有记录（只覆盖非空字段），返回是否有变化。"""
        if not keyword or not detail:
            return False
        current = (self.preset_details.get(keyword) or [{}])[0]
        merged = dict(current)
        for key, value in detail.items():
            if value not in (None, ""):
                merged[key] = value
        if merged == current:
            return False
        self.preset_details[keyword] = [merged]
        self.preset_details_dirty = True
        return True

    def catalog_detail_for(self, keyword, category):
        """从本地「全部课程」缓存里找出这门课，补齐课程名/教学班/时间/余量。

        缓存按教学班号索引（历次「拉取全部课程」的结果）。名称类预设
        （例如 GE:加工食品与健康生活）以前完全没有详情，列表里只剩一行裸名称；
        这里按 教学班号 → 课程号 → 课程名 依次回退匹配，并把命中的教学班数量一起带回。
        """
        catalog = getattr(self, "course_catalog", None) or {}
        if not catalog:
            return None, 0
        hit = catalog.get(keyword)
        if hit:
            return self._catalog_row_to_detail(keyword, hit), 1
        wanted = str(keyword or "").strip().casefold()
        if not wanted:
            return None, 0
        matches = [(teaching_class, row) for teaching_class, row in catalog.items()
                   if wanted in {str(row.get("course_code") or "").strip().casefold(),
                                 str(row.get("course_name") or "").strip().casefold()}]
        if not matches:
            return None, 0
        preferred = {"GE": "通识", "EN": "英语", "PE": "体育"}.get(category)
        if preferred:
            same_tab = [item for item in matches
                        if preferred in str(item[1].get("course_tab") or "")
                        or preferred in str(item[1].get("kclbmc") or "")]
            if same_tab:
                matches = same_tab
        teaching_class, row = matches[0]
        return self._catalog_row_to_detail(teaching_class, row), len(matches)

    @staticmethod
    def _catalog_row_to_detail(teaching_class, row):
        """把「全部课程」缓存的一行转成预设详情字段（与 course_preset_details.json 同构）。"""
        campus = row_campus(dict(row, teaching_class=teaching_class or row.get("teaching_class", "")))
        return {
            "course_name": row.get("course_name") or row.get("keyword") or "",
            "code": row.get("course_code") or row.get("kch") or "",
            "teaching_class": teaching_class or row.get("teaching_class") or "",
            "time": clean_time_text(row.get("time")),
            "room": clean_time_text(row.get("room")),
            "teacher": row.get("teacher") or "",
            "category": row.get("course_tab") or row.get("kclbmc") or "",
            "xf": row.get("xf", ""),                # 学分（凑学分/挑通识时看）
            "rongliang": row.get("rongliang", ""),
            "yixuan": row.get("yixuan", ""),
            "jxb_id": row.get("jxb_id", ""),
            "campus": campus,
            "platform": row_platform(dict(row, teaching_class=teaching_class or row.get("teaching_class", ""))),
        }

    @staticmethod
    def detail_from_fetch(found, jxb, category, jxb_count=1):
        """把一次在线查询到的教学班转成预设详情，供列表显示与落盘。"""
        course = getattr(found, "course", None)
        return {
            "course_name": getattr(course, "kcmc", "") or "",
            "code": getattr(course, "kch_id", "") or "",
            "teaching_class": getattr(jxb, "jxbmc", "") or "",
            "time": clean_time_text(getattr(jxb, "sksj", "")),
            "room": getattr(jxb, "jxdd", "") or "",
            "teacher": getattr(jxb, "jsxx", "") or "",
            "category": getattr(course, "tab_name", "") or category,
            "xf": getattr(jxb, "xf", ""),           # 学分
            "rongliang": getattr(jxb, "jxbrl", ""),
            "yixuan": getattr(jxb, "yxzrs", ""),
            "campus": detect_campus(getattr(jxb, "xqumc", ""), getattr(jxb, "jxbmc", ""),
                                    getattr(jxb, "jxdd", "")),
            "jxb_count": jxb_count,
        }

    def preset_detail_for(self, keyword, category):
        """取某条预设的详情；没有就先用本地缓存回填并记下来。"""
        details = self.preset_details.get(keyword)
        if details:
            return details[0]
        detail, count = self.catalog_detail_for(keyword, category)
        if not detail:
            return None
        if count > 1:
            detail["jxb_count"] = count
        detail["_local_only"] = True   # 来自本地缓存，人数可能过期
        self.preset_details[keyword] = [detail]
        # 列表控件里存的是展示行（类别｜课程名｜教学班｜…），二次刷新时它会被
        # parse_preset 还原成“教学班”而不是原来的课程名。这里同时按教学班建索引，
        # 保证两种键都能查到同一条详情，"刷新人数"的结果才贴得上去、显示才稳定。
        teaching_class = str(detail.get("teaching_class") or "").strip()
        if teaching_class and teaching_class != keyword and teaching_class not in self.preset_details:
            self.preset_details[teaching_class] = [detail]
        self.preset_details_dirty = True
        return detail

    def live_count_text(self, keyword, live_counts, detail):
        live = (live_counts or {}).get(keyword)
        if isinstance(live, str) and live:
            return live
        if isinstance(live, dict) and live.get("count"):
            return live["count"]
        cap, chosen = str(detail.get("rongliang") or ""), str(detail.get("yixuan") or "")
        if cap or chosen:
            suffix = "本地缓存" if detail.get("_local_only") else "上次查询"
            return f"余量 {cap or '?'} - 已选 {chosen or '?'}（{suffix}）"
        return "人数待实时刷新"

    def preset_display_text(self, raw, live_counts=None):
        """拼出待选列表的一整行：类别｜课程名｜关键字｜时间｜状态｜人数。

        关键约束：**第三列必须等于抢课时真正发给教务的查询关键字**。
        列表控件本身就是数据源（presets() 从控件里读），所以第三列写什么、
        抢课就会查什么。锁定教学班的条目显示教学班号；只写课程名的条目就显示
        课程名（服务端按名匹配、取第一个班），免得列表写着一个班、程序去抢另一个班。
        详情缺失时回退到本地「全部课程」缓存，两者都没有才原样显示。
        """
        keyword, category = parse_preset(raw)
        detail = self.preset_detail_for(keyword, category)
        if not detail:
            return raw
        count = self.live_count_text(keyword, live_counts, detail)
        tag = {"GE": "通识", "EN": "英语", "PE": "体育"}.get(category, "课程")
        class_name = str(detail.get("teaching_class") or "").strip()
        try:
            jxb_count = int(detail.get("jxb_count") or 1)
        except (TypeError, ValueError):
            jxb_count = 1
        subject = class_name if class_name and class_name == keyword else keyword
        time_text = clean_time_text(detail.get("time")) or "时间待补"
        if not looks_like_teaching_class(keyword):
            # 关键字是课程名 → 服务端按名匹配（多个教学班时取第一个班），
            # 在时间列如实标注，免得以为已经锁定了某个班。
            if class_name and jxb_count <= 1:
                extra = f"按课程名抢；缓存中唯一教学班 {class_name}"
            elif class_name:
                extra = f"按课程名抢；缓存中同名 {max(jxb_count, 1)} 个教学班，例如 {class_name}"
            else:
                extra = f"按课程名抢；缓存中同名 {max(jxb_count, 1)} 个教学班"
            time_text = f"{time_text}（{extra}）"
        campus = normalize_campus_name(detail.get("campus")) or detect_campus("", keyword, detail.get("room"))
        if is_off_campus(campus):
            # 非本校校区（教工路）：抢到了要跨校区通勤，列表里必须一眼看到。
            # 注意不能用「｜」（列表按它分段，第 3 列必须还是抢课关键字）。
            time_text = f"⚠{campus}校区（非你所在校区）· {time_text}"
        else:
            platform = str(detail.get("platform") or "").strip() or teaching_class_platform(keyword)
            if platform:
                # 网络课（智慧树/尔雅）：没有通勤问题，但也别以为是面授课。
                time_text = f"【{platform}】{time_text}"
        return f"{tag}｜{detail.get('course_name') or keyword}｜{subject}｜{time_text}｜待抢｜{count}"

    def refresh_preset_display(self, live_counts=None):
        current = self.presets() if hasattr(self, "preset_list") else []
        self.preset_list.delete(0, "end")
        for raw in current:
            self.preset_list.insert("end", self.preset_display_text(raw, live_counts))
        if getattr(self, "preset_details_dirty", False):
            self.save_preset_details()

    def _backup_presets(self):
        """分配/批量改写待选列表前留一份备份（只保留最近 5 份）。"""
        try:
            if not PRESET_FILE.exists():
                return
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            (DATA_DIR / f"course_presets.txt.bak_{stamp}").write_text(
                PRESET_FILE.read_text(encoding="utf-8-sig"), encoding="utf-8")
            for old in sorted(DATA_DIR.glob("course_presets.txt.bak_*"))[:-5]:
                old.unlink(missing_ok=True)
        except Exception as e:
            self.append(f"备份待选列表失败：{type(e).__name__}: {e}")

    def allocate_presets(self, apply=True, reason=""):
        """给每条预设分配具体教学班，避开时间段冲突（长期机制）。

        分配依据：已选课程占用的时间（教务已选 + 本程序抢到）+ 各类别还剩多少名额。
        结果会写回待选列表（改成锁定教学班）并写入详情表，日志里逐条说明原因。
        """
        raw_presets = [canonical_preset(raw) for raw in self.presets()]
        if not raw_presets:
            return [], ["[分配] 待选列表为空，无需分配。"]
        snapshot = self.progress_snapshot()
        done = {category: done_count
                for category, (done_count, _target) in snapshot.get("progress", {}).items()}
        catalog = self.course_catalog or {}
        selected = self.annotated_selected_courses()
        plan = plan_preset_allocation(raw_presets, self.occupied_times(), catalog,
                                      self.state_store.targets, done)
        report = [f"[分配] 触发原因：{reason or '手动'}"] + allocation_report(plan, selected)
        if not catalog:
            report.append("[分配] 本地还没有「全部课程」缓存，先在「全部课程」Tab 点一次「拉取全部课程」。")
        changed = 0
        new_presets = []
        for entry in plan:
            target = entry["raw"]
            if entry["status"] in {"已分配", "备用"} and entry.get("teaching_class"):
                category = entry["category"]
                locked = (f"{category}:{entry['teaching_class']}"
                          if category in {"GE", "EN", "PE"} else entry["teaching_class"])
                changed += 1 if locked != entry["raw"] else 0
                target = locked
                row = catalog.get(entry["teaching_class"])
                if row:
                    self.merge_preset_detail(entry["teaching_class"],
                                             self._catalog_row_to_detail(entry["teaching_class"], row))
            new_presets.append(target)
        if new_presets == raw_presets:
            # 没有任何变化就不要刷屏（例如每 60 秒对齐一次已选时的重复分配）
            note = "[分配] 无需调整：各条志愿的时间段与名额都没变化。"
            if self._last_allocation_report != note or reason == "手动点击":
                self.append(note)
            self._last_allocation_report = note
            if getattr(self, "preset_details_dirty", False):
                self.save_preset_details()
            return plan, [note]
        self._last_allocation_report = ""
        self._backup_presets()
        if apply:
            self.preset_list.delete(0, "end")
            for raw in new_presets:
                self.preset_list.insert("end", raw)
            self.persist_presets()
            if getattr(self, "preset_details_dirty", False):
                self.save_preset_details()
            self.refresh_preset_display()
            report.append(f"[分配] 已写回待选列表：{len(new_presets)} 条，其中 {changed} 条换成锁定的教学班。")
        for line in report:
            self.append(line)
        self.status.set("时间段分配完成")
        return plan, report

    def refresh_live_counts(self):
        """向教务实时查一遍每条预设的教学班/人数，并把结果写进详情表。"""
        base_url, cookie = self.base_var.get(), self.cookie_text()
        raw_presets = [canonical_preset(raw) for raw in self.presets()]
        def work():
            client = make_client(base_url, cookie, self.thread_log)
            result = {}
            for raw in raw_presets:
                keyword, category = parse_preset(raw)
                try:
                    found = fetch_course(client, keyword, "auto")
                    if found.jxbs:
                        j = found.jxbs[0]
                        if j.jxbrl or j.yxzrs:
                            text = f"余量 {j.jxbrl or '?'} - 已选 {j.yxzrs or '?'}（实时）"
                        else:
                            text = f"已找到{len(found.jxbs)}个教学班（实时）"
                        result[keyword] = {
                            "count": text,
                            "detail": self.detail_from_fetch(found, j, category, len(found.jxbs)),
                        }
                    else:
                        result[keyword] = {"count": "未找到/未开放"}
                except Exception as e:
                    result[keyword] = {"count": f"刷新失败：{type(e).__name__}"}
            return ("live_counts", result)
        self.run_thread(work, "刷新对应课程实时信息中……")

    def load_presets_from_file(self):
        """Load one course keyword/code per line from the local preset file."""
        try:
            if not PRESET_FILE.exists():
                return
            for raw in PRESET_FILE.read_text(encoding="utf-8-sig").splitlines():
                value = raw.strip()
                if value and not value.startswith("#"):
                    self.preset_list.insert("end", value)
        except Exception as e:
            self.append(f"读取预设文件失败：{type(e).__name__}: {e}")

    def start_monitor(self):
        try:
            interval = min(5.0, max(0.2, float(self.interval_var.get())))
        except Exception:
            interval = 0.3
        try:
            max_success = max(0, int(self.max_success_var.get()))
        except Exception:
            max_success = 1
        try:
            max_rounds = max(0, int(self.max_rounds_var.get()))
        except Exception:
            max_rounds = 0
        presets = self.presets()
        if not presets:
            messagebox.showwarning("缺预设", "请先添加预设课程/教学班。")
            return
        if self.avoid_conflict_var is not None and self.avoid_conflict_var.get():
            # 长期机制：每次开抢前按"已选课表 + 志愿名额"重新分配教学班，避开时间段冲突。
            self.allocate_presets(apply=True, reason="开始抢课前重新分配时间段")
            presets = self.presets()
        self.stop_event.clear()
        start_text = self.start_time_var.get().strip()
        base_url, cookie, selector = self.base_var.get(), self.cookie_text(), self.selector_var.get()
        retry, watch_full = bool(self.retry_var.get()), bool(self.watch_full_var.get())
        parallel, auto_after_init = bool(self.parallel_var.get()), bool(self.auto_after_init_var.get())
        # 门数一多，串行的每一轮大部分时间都花在"确认那些课还是满的"：
        # 优先盯前 N 门 → 只高频盯最想要的几门，其余保底每 3 轮试一次；0 = 盯全部。
        try:
            focus_count = max(0, int(float(self.focus_count_var.get() or 0)))
        except (TypeError, ValueError):
            focus_count = 0
        try:
            full_wait = min(FULL_WAIT_MAX, max(FULL_WAIT_MIN, float(self.full_wait_var.get())))
        except (TypeError, ValueError):
            full_wait = DEFAULT_FULL_WAIT
        self.save_runtime_settings()
        # 名额进度以"教务已选 + 本程序抢到"的并集为准，避免重复押注已满的类别
        # （例如教务里已有 1 门通识时，通识就只剩 1 个名额）。
        progress = self.progress_snapshot().get("progress", {})
        initial = {category: done for category, (done, _target) in progress.items()}
        # 教务已选 + 本程序已完成课程的**课程号**：同一门课的其它教学班不要再押
        # （教务会把先选上的替换掉，等于白费一次提交、还可能换成更差的班）。
        done_courses = sorted({str(course.get("course_code") or "").strip()
                               for course in self.system_selected
                               if str(course.get("course_code") or "").strip()}
                              | {str(record.get("course_code") or "").strip()
                                 for record in self.state_store.selected()
                                 if str(record.get("course_code") or "").strip()})
        self.run_thread(lambda: ("done", timed_monitor_rush(start_text, base_url, cookie, presets, selector, interval, max_success, max_rounds, retry, watch_full, parallel, auto_after_init, self.stop_event, self.thread_log, on_success=self.on_course_success, initial_category_success=initial, focus_count=focus_count, full_wait=full_wait, done_courses=done_courses)), "定时等待/监控抢课中……")

    def stop_monitor(self):
        self.stop_event.set()
        self.status.set("正在停止……")

    def search_course(self):
        kw = self.search_var.get().strip()
        mode = self.search_mode.get()
        # Tk 控件只能在主线程读：base_url / Cookie 先取出来，工作线程里只用这两个变量
        # （在工作线程里调 self.base_var.get() 会抛 "main thread is not in main loop"）。
        base_url, cookie = self.base_var.get(), self.cookie_text()
        self.save_search_settings()
        filters = json.loads(json.dumps(self.search_filters))
        refresh_availability = bool(getattr(self, "_refresh_availability_next", False))
        self._refresh_availability_next = False
        def work():
            rows = list(self.all_courses_rows or [])
            source = "本地缓存"
            if refresh_availability:
                try:
                    rows = pull_all_courses(base_url, cookie, self.thread_log)
                    if rows:
                        source = "在线刷新"
                        self.save_course_catalog(rows)
                    else:
                        self.thread_log("[余量刷新] 在线返回空数据，回退本地缓存")
                except Exception as e:
                    self.thread_log(f"[余量刷新] 在线访问失败，回退本地缓存：{type(e).__name__}: {e}")
            if not rows and COURSE_CATALOG_FILE.exists():
                try:
                    cached = json.loads(COURSE_CATALOG_FILE.read_text(encoding="utf-8"))
                    if isinstance(cached, dict):
                        rows = [dict(value, teaching_class=key) for key, value in cached.items()]
                except Exception as e:
                    self.thread_log(f"读取全部课程缓存失败：{type(e).__name__}: {e}")
            # 先按关键词命中全部，再依次按校区、周数筛，这样能如实报出各筛掉了几个。
            base_filters = dict(filters)
            base_filters["campus"] = CAMPUS_ALL
            base_filters["weeks"] = ""      # 周数在这一层单独算，好报出差值
            base_filters["credits"] = []    # 学分同理：先拿全量，再按学分收窄并报数
            hit_all = search_cached_courses(rows, kw, mode, base_filters)
            wanted_campus = filters.get("campus", CAMPUS_LOCAL)
            on_campus = [row for row in hit_all if campus_filter_match(row, wanted_campus)]
            hidden = len(hit_all) - len(on_campus)
            wanted_weeks = week_set(filters.get("weeks"))
            weeks_mode = filters.get("weeks_mode") or WEEK_MODE_INCLUDE
            if wanted_weeks:
                found = [row for row in on_campus if week_filter_match(
                    row.get("time") or row.get("sksj") or "", wanted_weeks, weeks_mode)]
                weeks_hidden = len(on_campus) - len(found)
            else:
                found, weeks_hidden = on_campus, 0
            wanted_credits = [str(x) for x in (filters.get("credits") or [])]
            if wanted_credits:
                cred_found = [row for row in found if credit_matches(row, wanted_credits)]
                credits_hidden = len(found) - len(cred_found)
                found = cred_found
                self.thread_log(f"[本地搜索] 学分「{'/'.join(wanted_credits)} 分」隐藏 {credits_hidden} 个"
                                "（含学分未知的）")
            else:
                credits_hidden = 0
            self.thread_log(f"[本地搜索] 关键词「{kw}」命中 {len(hit_all)} 个教学班"
                            + (f"，按校区「{wanted_campus}」隐藏 {hidden} 个（教工路等非本校校区）"
                               if hidden else "")
                            + (f"，按周数「{WEEK_MODE_BUTTON_TEXT.get(weeks_mode, weeks_mode)} "
                               f"{filters.get('weeks')}」隐藏 {weeks_hidden} 个（含时间/周次识别不出的）"
                               if wanted_weeks and weeks_hidden else ""))
            return ("search_local", LocalSearchResult(rows=found, keyword=kw, source=source,
                                                      hidden=hidden, campus=str(wanted_campus or ""),
                                                      weeks_hidden=weeks_hidden,
                                                      weeks_cond=str(filters.get("weeks") or ""),
                                                      weeks_mode=weeks_mode if wanted_weeks else "",
                                                      credits_hidden=credits_hidden,
                                                      credits_cond="/".join(wanted_credits)))
        self.run_thread(work, "搜索课程中……")

    def add_search_to_preset(self):
        """把搜索里**勾选的多行**一起加入自动抢课预设。

        默认锁定到「教学班」（即你在搜索结果里勾选的那一行）：
        1) 抢课时查的就是这个班，不会再像按课程名那样被匹配到别的 Tab
           （例如「数字人文」被匹配成主修课程里的别班）；
        2) 课程名/时间/教室/教师/人数/校区一起写进详情表，待选列表才是完整的一行。
        想按课程名加入（哪一班都能抢）时，把搜索框右侧的模式切成「课程名」。
        """
        result = getattr(self, "search_result", None)
        if not result:
            messagebox.showwarning("未搜索", "请先搜索课程。")
            return
        indices = self.checked_search_indices()
        if not indices:
            messagebox.showwarning("未勾选", "请先勾选要加入的教学班（点列表第一列的方框，或用「☑ 全选」）。")
            return
        mode = self.search_mode.get()
        existing = {parse_preset(self.preset_list.get(i))[0] for i in range(self.preset_list.size())}
        added, skipped, off_campus = [], [], []
        for index in indices:
            found = self._search_entry(index)
            if not found:
                continue
            keyword, category, detail = self._preset_from_search_row(found[0], found[1], mode, result)
            if not keyword:
                continue
            if keyword in existing:
                skipped.append(keyword)
                continue
            prefix = {"体育": "PE", "英语": "EN", "通识": "GE"}.get(category)
            self.preset_list.insert("end", f"{prefix}:{keyword}" if prefix else keyword)
            existing.add(keyword)
            if detail:
                self.merge_preset_detail(keyword, detail)
            added.append(keyword)
            if is_off_campus((detail or {}).get("campus")):
                off_campus.append(f"{detail.get('course_name') or keyword}（{detail.get('campus')}）")
        if not added:
            if skipped:
                messagebox.showinfo("已在列表", f"勾选的 {len(skipped)} 门课都已在自动抢课预设列表中。")
            return
        self.save_preset_details()
        self.persist_presets()
        self.refresh_preset_display()
        self.update_pick_hint()
        summary = "、".join(added[:6]) + ("…" if len(added) > 6 else "")
        self.append(f"已加入 {len(added)} 门课程：{summary}"
                    + (f"（{len(skipped)} 门已在列表中，已跳过）" if skipped else "")
                    + "（可在「自动抢课」Tab 调整顺序后开始监控）")
        self.status.set(f"已加入 {len(added)} 门课程")
        self.nb.select(self.auto_tab)
        if self.avoid_conflict_var is not None and self.avoid_conflict_var.get():
            # 加入后立刻按"已选课表 + 名额"分配时间段：这一步也会把最详细的信息
            # （课程名/教学班/时间/教室/教师/余量）写进详情表。
            self.allocate_presets(apply=True, reason="加入新课程后自动分配时间段")
        if off_campus:
            # 教工路校区单程 40 分钟以上：抢到了才发现要跨校区就晚了，必须显式提醒。
            self.append("⚠ 其中 %d 门不在你所在校区（下沙）：%s —— 教工路校区需跨校区通勤；"
                        "不想要就删掉这几条，或把「上课校区」筛成「下沙」。" % (len(off_campus), "、".join(off_campus)))
            messagebox.showwarning("注意：非本校校区课程",
                                   "以下课程不在你所在校区（下沙）：\n"
                                   + "\n".join(f"· {name}" for name in off_campus)
                                   + "\n\n教工路校区需要跨校区通勤；不想抢的话，"
                                     "请在「自动抢课」Tab 删除这几条（待选列表里已用 ⚠教工路 标出）。")

    def submit_search_selected(self):
        if not self.search_result:
            messagebox.showwarning("未搜索", "请先搜索课程。")
            return
        indices = self.checked_search_indices()
        if not indices:
            messagebox.showwarning("未勾选", "请先勾选要提交的教学班（点列表第一列的方框，或用「☑ 全选」）。")
            return
        targets = []
        for index in indices:
            found = self._search_entry(index)
            if not found:
                continue
            row, jxb = found
            if row is not None:
                label = row.get("teaching_class") or row.get("course_name") or ""
                time_text = row.get("time") or ""
            else:
                label, time_text = getattr(jxb, "jxbmc", ""), getattr(jxb, "sksj", "")
            targets.append((index, row, jxb, f"{label}  {time_text}".strip()))
        if not targets:
            messagebox.showwarning("未勾选", "勾选的行已失效，请重新搜索。")
            return
        preview = "\n".join(f"· {item[3]}" for item in targets[:8])
        if len(targets) > 8:
            preview += f"\n… 共 {len(targets)} 个教学班"
        if not messagebox.askyesno("确认批量提交", f"确认同时提交这 {len(targets)} 个教学班？\n{preview}"):
            return
        loop_enabled = bool(self.search_loop_var.get())
        pending = [canonical_preset(raw) for raw in self.presets()]
        result = self.search_result
        # base_url / Cookie 必须在主线程取好；工作线程里读 Tk 控件会抛
        # "main thread is not in main loop"（旧版这里就是这么挂的）。
        base_url, cookie = self.base_var.get(), self.cookie_text()
        off_campus = [item[3] for item in targets
                      if is_off_campus(row_campus(item[1] if item[1] is not None
                                                  else {"teaching_class": getattr(item[2], "jxbmc", ""),
                                                        "xqumc": getattr(item[2], "xqumc", "")}))]
        if off_campus:
            self.thread_log("⚠ 本次提交里包含非本校校区（教工路）的教学班：" + "、".join(off_campus[:3]))

        def work():
            shared = {"client": None}

            def get_client():
                """本地行要先去教务查一次教学班；整批共用一个连接，只登录一次。"""
                if shared["client"] is None:
                    shared["client"] = make_client(base_url, cookie, self.thread_log)
                    init_until_open(shared["client"], self.thread_log, once=True)
                return shared["client"]

            succeeded, failed = [], []
            queue = list(targets)
            attempt = 0
            while queue:
                attempt += 1
                retry = []
                for index, row, jxb, label in queue:
                    if self.stop_event.is_set():
                        return ("status", f"批量抢课已停止：成功 {len(succeeded)} 个，失败 {len(failed)} 个")
                    try:
                        if row is not None:
                            client = get_client()
                            found = fetch_course(client, row.get("teaching_class") or row.get("kch", ""), "教学班")
                            if not found.jxbs:
                                raise RuntimeError("该课程对应的教学班当前无法获取")
                            wanted_id = row.get("jxb_id") or row.get("do_id")
                            picked = next((item for item in found.jxbs
                                           if wanted_id and wanted_id in {getattr(item, "jxb_id", ""),
                                                                          getattr(item, "do_id", "")}),
                                          found.jxbs[0])
                            target_result = found
                        else:
                            client, target_result, picked = result.client, result.course, jxb
                        flag, msg = select_jxb(client, target_result, picked, self.thread_log)
                    except Exception as e:
                        failed.append((label, f"{type(e).__name__}: {e}"))
                        self.thread_log(f"批量抢课「{label}」异常：{type(e).__name__}: {e}")
                        continue
                    if str(flag) == "1":
                        actual_class = getattr(picked, "jxbmc", "") or label.split("  ")[0]
                        matched = match_pending_preset(pending,
                                                       getattr(target_result, "keyword", "") or result.keyword,
                                                       actual_class)
                        self.on_course_success(matched, {
                            "teaching_class": actual_class,
                            "teacher": getattr(picked, "jsxx", ""), "jxb_id": getattr(picked, "jxb_id", ""),
                            "kch_id": getattr(target_result.course, "kch_id", ""),
                        })
                        succeeded.append(label)
                        self.thread_log(f"✔ 已选上：{label}")
                        continue
                    if loop_enabled:
                        retry.append((index, row, jxb, label))
                    else:
                        failed.append((label, describe_select_result(flag, msg)))
                        self.thread_log(f"✘ {label}：{describe_select_result(flag, msg)}")
                if not loop_enabled or not retry:
                    break
                queue = retry
                if self.stop_event.is_set():
                    break
                self.thread_log(f"批量抢课第 {attempt} 轮未全部成功（剩 {len(retry)} 个），1.5 秒后重试")
                time.sleep(1.5)
            for label, why in failed[:8]:
                self.thread_log(f"[批量抢课] 未成功：{label} → {why}")
            summary = f"批量抢课结束：成功 {len(succeeded)} 个，失败 {len(failed)} 个"
            if succeeded:
                summary += "；已选上：" + "、".join(succeeded[:3])
            return ("status", summary)

        self.stop_event.clear()
        self.run_thread(work, f"批量提交中（{len(targets)} 个教学班）……")

    def export_all_courses_excel(self):
        rows = getattr(self, "all_courses_rows", []) or []
        if not rows:
            messagebox.showwarning("没有数据", "请先拉取全部课程，再导出 Excel。")
            return
        default_name = "浙江工商全部课程.xlsx"
        path = filedialog.asksaveasfilename(
            title="导出全部课程 Excel",
            defaultextension=".xlsx",
            initialfile=default_name,
            initialdir=str(DATA_DIR),
            filetypes=[("Excel 工作簿", "*.xlsx")],
        )
        if not path:
            return
        try:
            export_courses_to_xlsx(rows, path)
            self.append(f"已导出 Excel：{path}")
            self.status.set("Excel 导出完成")
            messagebox.showinfo("导出完成", "已保存：\n" + path)
        except Exception as e:
            messagebox.showerror("导出失败", f"{type(e).__name__}: {e}")

    def pull_all_courses_ui(self):
        base_url, cookie = self.base_var.get(), self.cookie_text()
        def work():
            rows = pull_all_courses(base_url, cookie, self.thread_log)
            return ("all_courses", rows)
        self.run_thread(work, "拉取全部课程中……")

    def poll(self):
        try:
            while True:
                kind, payload = self.q.get_nowait()
                if kind == "log":
                    self.update_task_progress(payload)
                    self.append(payload)
                elif kind == "cookie":
                    self.cookie_box.delete("1.0", "end")
                    self.cookie_box.insert("1.0", payload)
                    self.status.set("Cookie 已填入")
                    self.append("Cookie 已自动填入。")
                    self.finish_task_progress("Cookie 已填入")
                    self.after(200, self.refresh_selected_courses)
                elif kind == "status":
                    self.status.set(payload)
                    self.finish_task_progress("已完成")
                    self.append(payload)
                elif kind == "live_counts":
                    counts, details_changed = {}, False
                    for keyword, entry in (payload or {}).items():
                        if isinstance(entry, dict):
                            counts[keyword] = entry.get("count", "")
                            if entry.get("detail") and self.merge_preset_detail(keyword, entry["detail"]):
                                details_changed = True
                        else:
                            counts[keyword] = entry
                    if details_changed:
                        self.save_preset_details()
                    self.refresh_preset_display(counts)
                    self.status.set("对应课程实时信息已刷新")
                    self.append("已刷新预设课程：本次显示为教务系统实时查询结果。")
                    self.finish_task_progress("已刷新人数")
                elif kind == "search":
                    self.finish_task_progress("搜索完成")
                    self.search_result = payload
                    entries = []
                    for j in payload.jxbs:
                        entries.append({
                            "name": getattr(payload.course, "kcmc", "") or payload.keyword,
                            "class": getattr(j, "jxbmc", "") or getattr(j, "jxb_id", ""),
                            "credit": credit_value({"xf": getattr(j, "xf", "")}),
                            "time": clean_time_text(getattr(j, "sksj", "")),
                            "teacher": getattr(j, "jsxx", ""),
                            "status": "可提交",
                            "campus": detect_campus(getattr(j, "xqumc", ""),
                                                    getattr(j, "jxbmc", ""),
                                                    getattr(j, "jxdd", "")),
                            "platform": teaching_class_platform(getattr(j, "jxbmc", "")),
                        })
                    off = self.fill_search_tree(entries)
                    self.status.set(f"搜索完成：{len(payload.jxbs)} 个教学班"
                                    + (f"（{off} 个在非本校校区，已红色标注）" if off else ""))
                elif kind == "search_local":
                    self.finish_task_progress("搜索完成")
                    self.search_result = payload
                    entries = []
                    for row in payload.rows:
                        entries.append({
                            "name": row.get("course_name", "") or row.get("keyword", ""),
                            "class": row.get("teaching_class", "") or row.get("jxb_id", ""),
                            "credit": credit_value(row),
                            # 教务的时间字段带 <br/>，直接塞进表格会露出原始标签
                            # （真实数据里「丝绸之路与中外文化交流」就是这样），统一清洗。
                            "time": clean_time_text(row.get("time", "")),
                            "teacher": row.get("teacher", ""),
                            "status": "可提交" if row.get("do_id") else "需开放后提交",
                            "campus": row.get("campus") or row_campus(row),
                            "platform": row_platform(row),
                        })
                    off = self.fill_search_tree(entries)
                    note = ""
                    if off:
                        note = f"，其中 {off} 个在非本校校区（已红色标注 ⚠）"
                    if getattr(payload, "hidden", 0):
                        note += f"，另有 {payload.hidden} 个被「上课校区={payload.campus}」筛掉"
                    if getattr(payload, "weeks_hidden", 0):
                        note += (f"，另有 {payload.weeks_hidden} 个被「周数{payload.weeks_mode}"
                                 f"={payload.weeks_cond}」筛掉（含时间/周次识别不出的）")
                    if getattr(payload, "credits_hidden", 0):
                        note += (f"，另有 {payload.credits_hidden} 个被「学分={payload.credits_cond}"
                                 " 分」筛掉（含学分未知的）")
                    self.status.set(f"搜索完成：{len(payload.rows)} 个教学班（{payload.source}）{note}")
                elif kind == "all_courses":
                    self.all_courses_rows = payload
                    self.save_course_catalog(payload)
                    self.rebuild_credit_buttons()    # 新数据可能带来新的学分档位
                    for r in self.all_tree.get_children():
                        self.all_tree.delete(r)
                    for i, row in enumerate(payload, 1):
                        cap = row.get("rongliang", "")
                        yx = row.get("yixuan", "")
                        sel = f"{yx}/{cap}" if (yx != "" or cap != "") else ""
                        campus = row.get("campus") or row_campus(row)
                        platform = row_platform(row)
                        campus_text = (f"⚠{campus}" if is_off_campus(campus)
                                       else "网课" if platform else campus)
                        self.all_tree.insert("", "end", values=(
                            i,
                            row.get("course_name", "") or row.get("keyword", ""),
                            row.get("kch", ""),
                            row.get("teaching_class", ""),
                            campus_text,
                            (f"【{platform}】" if platform else "") + str(row.get("time", "")),
                            row.get("room", ""),
                            sel,
                        ), tags=("offcampus",) if is_off_campus(campus) else ())
                    self.status.set(f"全部课程拉取完成：{len(payload)} 个教学班，可导出 Excel")
                    self.finish_task_progress("拉取全部课程完成")
                elif kind == "selected_courses":
                    if int(payload.get("count", 0)) == 0 and self.system_selected_count > 0:
                        self.status.set("已选课程刷新异常：保留旧缓存")
                        self.append("已选接口本次返回空列表，已保留旧缓存，未覆盖现有已选课程。")
                        self.last_selected_refresh = datetime.now()   # 避免每 5 秒重试打爆接口
                        continue
                    self.system_selected = payload.get("courses", [])
                    self.system_selected_count = int(payload.get("count", len(self.system_selected)))
                    self.system_selected_credits = str(payload.get("credits", ""))
                    self.last_selected_refresh = datetime.now()
                    self.save_system_selected_cache(payload)
                    moved = self.state_store.reconcile_system_selected(self.system_selected)
                    if moved:
                        self.reload_pending_presets()
                        self.append("已从待选列表移除教务系统已选课程：" + "、".join(moved))
                    # 反向：教务已选里已经没有、但本程序记着“已选中”的课程 → 放回待选
                    # （志愿没抽中/被退课都属于这种，列表要能自动回到“待抢”状态）
                    system_classes = {c.get("teaching_class") for c in self.system_selected if c.get("teaching_class")}
                    system_names = {str(c.get("course_name") or "").strip() for c in self.system_selected}
                    missing = [record for record in self.state_store.selected()
                               if str(record.get("teaching_class") or "").strip()
                               and record.get("teaching_class") not in system_classes
                               and str(record.get("course_name") or "").strip() not in system_names]
                    restored = self.state_store.demote_to_pending(missing)
                    if restored:
                        self.reload_pending_presets()
                        self.append("教务已选里已不存在、已放回待选列表：" + "、".join(restored))
                    self.update_selected_dashboard()
                    self.status.set(f"已选课程已刷新：{self.system_selected_count} 门")
                    self.append(f"已选课程刷新完成：{self.system_selected_count} 门，{self.system_selected_credits or '?'} 学分。")
                    # 已选课程有变动（新增/被退）→ 重新对齐时间段分配
                    signature = tuple(sorted((course.get("teaching_class") or "", str(course.get("status") or ""))
                                             for course in self.system_selected))
                    if signature != getattr(self, "_selected_signature", None):
                        self._selected_signature = signature
                        if self.avoid_conflict_var is not None and self.avoid_conflict_var.get():
                            self.allocate_presets(apply=True, reason="已选课程有变动，重新对齐时间段")
                    self.finish_task_progress("已刷新已选课程")
                elif kind == "course_success":
                    raw_preset, _local_snapshot = payload
                    self.reload_pending_presets()
                    # 进度用"教务已选 + 本程序抢到"的并集，抢到的立刻计入、并从待选移除
                    snapshot = self.progress_snapshot()
                    self.update_selected_dashboard(snapshot)
                    keyword, _category = parse_preset(raw_preset)
                    self.status.set("选课完成" if snapshot.get("completed") else f"已选中：{keyword}")
                    self.append(f"成功课程已移入“已选课程”并从待选列表剔除：{keyword}")
                elif kind == "done":
                    self.finish_task_progress("已完成")
                    if isinstance(payload, dict):
                        snapshot = self.progress_snapshot()
                        self.update_selected_dashboard(snapshot)
                        done_all = bool(payload.get("completed") or snapshot.get("completed"))
                        self.status.set("选课完成" if done_all else "监控结束，选课尚未完成")
                    else:
                        self.status.set("任务结束")
                    # 抢课一结束就用教务真相对齐一次进度
                    self.request_selected_refresh_soon()
                elif kind == "error":
                    self.finish_task_progress("出错")
                    self.status.set("出错")
                    self.append("[错误]")
                    self.append(payload)
                    first = payload.split("\n", 1)[0]
                    if "NotyetStarted" in payload:
                        self.status.set("选课未开放：拉取全部课程需等开放后使用")
                        self.append("提示：NotyetStarted 表示选课尚未开放，此时无法拉取课程列表。开放后再点一次即可。")
                        messagebox.showinfo("选课未开放", "选课尚未开放，暂时无法拉取全部课程。\n\n等选课开放后再点一次「拉取全部课程」即可。")
                    else:
                        messagebox.showerror("错误", first)
        except queue.Empty:
            pass
        self.settle_task_progress()
        try:
            self._poll_after_id = self.after(100, self.poll)
        except tk.TclError:
            self._poll_after_id = None


if __name__ == "__main__":
    App().mainloop()
